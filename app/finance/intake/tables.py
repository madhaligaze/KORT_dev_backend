"""Файл → таблицы: один читатель для всех входов приёмника.

До приёмника у каждого разбора было своё чтение: реестр читал xlsx полным
openpyxl ради ширин и заливки, выписки - быстрым построчным и выбирали одну
вкладку, сводная приходила строками Google. Здесь чтение одно, а что
сохранить - решает вызывающий:

* `detail=True` - всё, что говорит о смысле строк, кроме значений: скрытые
  строки (черновик, который не заключили, - его нельзя завести молча),
  объединённые ячейки (шапка в две строки: «Договор» над «№ | дата»),
  ширины колонок и заливка шапки. Это даёт только полный openpyxl.
* `detail=False` - только значения, построчным чтением: выписка на тысячи
  строк читается так же быстро, как раньше.

Формат решает содержимое (`formats.sniff`), а не расширение. Значения ячеек
не приводятся к тексту: дата, пришедшая датой, не имеет неоднозначности
порядка частей.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any, Sequence

from app.finance import formats

#: Сколько строк и колонок листа читаем. Больше - не таблица учёта, а выгрузка
#: чего-то другого; реестр BBC - 600 строк, журнал денег - 6 000.
MAX_ROWS = 20_000
MAX_COLS = 120
#: Заливка нужна для шапок - дальше этой строки её не собираем.
FILL_ROWS = 400


class TableError(RuntimeError):
    """Файл не читается таблицей - с текстом для человека."""


@dataclass
class Table:
    """Лист: значения и то, что о строках говорит оформление.

    Номера строк и колонок - с нуля, как в `rows`. Человеку строка
    показывается номером Excel (`line + 1`).
    """

    name: str
    rows: list[list[Any]]
    #: Строки, спрятанные в источнике.
    hidden: frozenset[int] = frozenset()
    #: Объединённые ячейки: (верх, лево, низ, право) включительно.
    merged: tuple[tuple[int, int, int, int], ...] = ()
    widths: dict[int, float] = field(default_factory=dict)
    fills: dict[tuple[int, int], str] = field(default_factory=dict)
    #: Уровень строки: отступ первой ячейки или уровень группировки строк.
    #: В ОСВ 1С так стоят уровни «контрагент → договор»; нет записи - уровень 0.
    levels: dict[int, int] = field(default_factory=dict)

    def owner(self, row: int, col: int) -> tuple[int, int] | None:
        """Левая верхняя ячейка объединения, в которое входит ячейка; `None` - не входит."""
        for top, left, bottom, right in self.merged:
            if top <= row <= bottom and left <= col <= right:
                return top, left
        return None

    def spans(self, row: int) -> list[tuple[int, int]]:
        """Объединения, которые начинаются в строке и идут вправо: (лево, право)."""
        return [(left, right) for top, left, _bottom, right in self.merged if top == row and right > left]


def _color(fill: Any) -> str:
    try:
        rgb = fill.start_color.rgb if fill is not None and fill.fill_type else None
    except AttributeError:
        return ""
    return rgb[-6:] if isinstance(rgb, str) and len(rgb) >= 6 and rgb not in ("00000000",) else ""


def read(
    data: bytes, file_name: str, *, detail: bool = True, max_rows: int | None = MAX_ROWS, hidden_sheets: bool = False
) -> list[Table]:
    """Все видимые листы файла. Скрытый лист не читается: его не видит и человек.

    `hidden_sheets=True` и `max_rows=None` - как читали выписки до приёмника:
    все листы, без предела строк (годовая выписка Kaspi Gold - 2 000 строк,
    журнал денег - 6 000, и лист с шапкой выбирает сам разбор).
    """
    kind = formats.sniff(data, file_name)
    if kind == "xlsx":
        return _read_xlsx(data, detail=detail, max_rows=max_rows, hidden_sheets=hidden_sheets)
    if kind == "pdf":
        raise TableError("Это PDF - он читается разбором выписок, а не как таблица")
    if kind == "image":
        raise TableError(
            "Это изображение. Загрузите выписку из банка в PDF или Excel - "
            "снимок экрана не содержит всех строк и реквизитов."
        )
    if kind == "text" and b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise TableError(
            f"Не понимаю формат файла «{file_name}». Подойдут PDF, Excel (.xlsx, .xls), "
            "CSV и выгрузка банк-клиента для 1С."
        )
    try:
        tables = formats.read_tables(data, kind)
    except formats.FormatError as exc:
        raise TableError(str(exc)) from exc
    stem = file_name.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
    return [Table(stem if len(tables) == 1 else f"{stem} #{index + 1}", [list(row) for row in rows]) for index, rows in enumerate(tables)]


def _read_xlsx(data: bytes, *, detail: bool, max_rows: int | None, hidden_sheets: bool = False) -> list[Table]:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - библиотека в зависимостях
        raise TableError("На сервере нет openpyxl - чтение Excel недоступно") from exc
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=not detail)
    except Exception as exc:  # noqa: BLE001 - сюда попадает любой битый zip
        raise TableError("Файл не читается как Excel (.xlsx)") from exc
    out: list[Table] = []
    try:
        for sheet in book.worksheets:
            if not hidden_sheets and getattr(sheet, "sheet_state", "visible") != "visible":
                continue
            out.append(_sheet(sheet, detail=detail, max_rows=max_rows))
    finally:
        book.close()
    return out


def _sheet(sheet: Any, *, detail: bool, max_rows: int | None) -> Table:
    if not detail:
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if max_rows is not None and index >= max_rows:
                raise TableError(f"Лист «{sheet.title}»: больше {max_rows} строк")
            rows.append(list(row))
        return Table(sheet.title, rows)

    from openpyxl.utils import column_index_from_string

    rows = []
    fills: dict[tuple[int, int], str] = {}
    levels: dict[int, int] = {}
    max_col = min(sheet.max_column or 0, MAX_COLS)
    for r_index, row in enumerate(sheet.iter_rows(max_col=max_col)):
        if max_rows is not None and r_index >= max_rows:
            raise TableError(f"Лист «{sheet.title}»: больше {max_rows} строк")
        values = []
        for c_index, cell in enumerate(row):
            values.append(cell.value)
            if cell.value not in (None, "") and r_index < FILL_ROWS:
                color = _color(cell.fill)
                if color:
                    fills[(r_index, c_index)] = color
        if row and row[0].value not in (None, ""):
            indent = getattr(getattr(row[0], "alignment", None), "indent", 0) or 0
            if indent:
                levels[r_index] = int(indent)
        rows.append(values)
    hidden = frozenset(
        number - 1 for number, dimension in sheet.row_dimensions.items() if dimension.hidden and 0 < number <= len(rows)
    )
    for number, dimension in sheet.row_dimensions.items():
        outline = getattr(dimension, "outline_level", 0) or 0
        if outline and 0 < number <= len(rows):
            levels[number - 1] = max(levels.get(number - 1, 0), int(outline))
    merged = tuple(
        (rng.min_row - 1, rng.min_col - 1, rng.max_row - 1, rng.max_col - 1)
        for rng in sheet.merged_cells.ranges
        if rng.min_row - 1 < len(rows)
    )
    widths: dict[int, float] = {}
    for letter, dimension in sheet.column_dimensions.items():
        try:
            index = column_index_from_string(letter) - 1
        except ValueError:
            continue
        if dimension.width:
            widths[index] = round(float(dimension.width), 1)
    return Table(sheet.title, rows, hidden, merged, widths, fills, levels)


def from_values(name: str, values: Sequence[Sequence[Any]]) -> Table:
    """Строки, прочитанные не из файла (лист Google), - та же таблица."""
    return Table(name, [list(row) for row in values])


__all__ = ["MAX_COLS", "MAX_ROWS", "Table", "TableError", "from_values", "read"]
