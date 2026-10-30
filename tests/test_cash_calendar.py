"""Платёжный календарь ядра - компания руками, из таблицы и из журнала.

Проверяется то, ради чего календарь в ядре, а не в формулах таблицы:
* остаток на утро - факт счетов юрлица, дни до сегодня - факт, дальше - ожидания;
* поступление по договору ждётся в срок месяца на рабочий день (праздники
  страны), долг без обещания - без даты, обещание ставит долг на день;
* плановый расход сводится с фактом, прошедший неоплаченный - на сегодня;
* разрыв у юрлица - с подсказкой перевести с юрлица, у которого есть запас;
* сценарии и строки, не идущие в остаток; «Отменить» возвращает;
* лист файлом: формы «платёж по плану» и «обещание по долгу» по названиям
  колонок, повторная загрузка не двоит, лист - правда для своих строк;
* повторяющийся расход журнала - предложение, принятое - плановый платёж.
"""
from __future__ import annotations

import base64
import io
from datetime import date

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.finance import audit, auth
from app.finance.contracts import service

BASE = "/api/v1/finance"
MODEL = f"{BASE}/model"
PASSWORD = "pass-12345"
CLIENT_BIN = "260540028486"
TODAY = date(2026, 10, 27)


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.core.config import settings
    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    monkeypatch.setattr(settings, "environment", "test")
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()
    yield engine
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()


@pytest.fixture
def owner(finance_db, monkeypatch) -> TestClient:
    from finance_routes import finance_app

    from app.finance.model import ask as model_ask

    monkeypatch.setattr(service, "today", lambda: TODAY)
    monkeypatch.setattr(model_ask, "today", lambda: TODAY)
    person = TestClient(finance_app(), headers={"x-client-ip": "10.0.0.8", "user-agent": "pytest-agent"})
    response = person.post(
        f"{BASE}/auth/register",
        json={"email": "owner@kassa.kz", "password": PASSWORD, "company": "Касса (тест)", "full_name": "Сериков Арман"},
    )
    assert response.status_code == 201, response.text
    return person


def _ok(response, status: int = 200) -> dict:
    assert response.status_code == status, response.text
    return response.json()


def _apply(client: TestClient, *commands: dict) -> dict:
    return _ok(client.post(f"{MODEL}/apply", json={"commands": list(commands)}))


def _one(client: TestClient, command: str, **params) -> dict:
    return _apply(client, {"command": command, "params": params})["commands"][0]["result"]


def _calendar(client: TestClient, **params) -> dict:
    return _ok(client.post(f"{MODEL}/ask", json={"source": "calendar", "params": params}))


def _entity(calendar: dict, code: str) -> dict:
    return next(item for item in calendar["entities"] if item["title"] == code)


def _value(calendar: dict, code: str, key: str, day: str) -> str | None:
    board = _entity(calendar, code)
    line = next((item for item in board["lines"] if item["key"] == key), None)
    return line["values"][calendar["days"].index(day)] if line else None


def _closing(calendar: dict, code: str, day: str) -> str:
    return _entity(calendar, code)["closing"][calendar["days"].index(day)]


def _company(client: TestClient) -> dict:
    alpha = _one(client, "entity.create", name="ТОО Альфа", code="АЛФ")
    beta = _one(client, "entity.create", name="ТОО Бета", code="БЕТ")
    _one(client, "accounting.start", since="2026-09-01")
    party = _one(client, "participant.create", name='ТОО "Клиент"', bin=CLIENT_BIN)
    made = _one(
        client, "contract.create",
        values={"executor": {"id": alpha["id"]}, "customer": {"id": party["id"]}, "number": "К-1", "amount": "100000",
                "billing": "month", "signed_at": "2026-09-03"},
    )
    contract = made["contract"]["id"]
    _one(client, "payment.record", day="2026-09-05", amount="100000", contract=contract, month="2026-09", key="k1-sep")
    _one(client, "payment.record", day="2026-10-01", amount="300000", entity=alpha["id"], party="Учредитель", key="loan")
    _one(client, "payment.record", day="2026-10-01", amount="500000", entity=beta["id"], party="Учредитель", key="beta")
    return {"alpha": alpha["id"], "beta": beta["id"], "contract": contract}


def test_kalendar_rukami_srok_obeshanie_plan_razryv(owner: TestClient) -> None:
    made = _company(owner)
    # Факт расхода 20.10 и плановые платежи: один сведётся с фактом, второй - перенесётся на сегодня.
    _one(owner, "payment.record", direction="out", day="2026-10-20", amount="50000", entity=made["alpha"], party="Провайдер", key="net")
    _one(owner, "cash.plan", direction="out", entity=made["alpha"], day="2026-10-19", amount="50000", line="Интернет")
    _one(owner, "cash.plan", direction="out", entity=made["alpha"], day="2026-10-21", amount="70000", line="Курьер")

    calendar = _calendar(owner)
    assert calendar["from"] == "2026-10-26" and calendar["to"] == "2026-11-29", "окно - с понедельника пять недель"
    assert "2026-10-26" in calendar["days_off"], "25.10 - воскресенье: праздник переносится на понедельник (пакет РК)"
    alpha = _entity(calendar, "АЛФ")
    assert alpha["opening"][0] == "350000.00", "остаток на утро - факт: 100 000 + 300 000 - 50 000"
    assert _entity(calendar, "БЕТ")["opening"][0] == "500000.00"
    # Октябрь (срок 03.10) не оплачен - долг без обещания: без даты, в остаток не идёт.
    debt_line = next(item for item in alpha["lines"] if item["key"] == "in:debt:договоры")
    assert debt_line["no_date"] == "100000.00"
    # Ноябрь - в срок 03.11 (вторник).
    assert _value(calendar, "АЛФ", "in:schedule:договоры", "2026-11-03") == "100000.00"
    # Сведённый план - сделан; несведённый прошедший - на сегодня.
    assert any(item["how"] == "сведено" for item in calendar["done"])
    carried = [item for item in calendar["items"] if item["state"] == "carried"]
    assert [(item["day"], item["amount"]) for item in carried] == [("2026-10-27", "70000.00")]
    assert _closing(calendar, "АЛФ", "2026-10-27") == "280000.00"

    promise = _apply(owner, {"command": "cash.promise", "params": {"contract": made["contract"], "day": "2026-10-29", "note": "обещал в четверг"}})
    _one(owner, "cash.plan", direction="out", entity=made["alpha"], day="2026-10-30", amount="650000", line="Аренда офисов")
    calendar = _calendar(owner)
    assert _value(calendar, "АЛФ", "in:debt:договоры", "2026-10-29") == "100000.00", "обещание ставит долг на день"
    assert _closing(calendar, "АЛФ", "2026-10-29") == "380000.00"
    assert _closing(calendar, "АЛФ", "2026-10-30") == "-270000.00"
    gap = next(item for item in calendar["gaps"] if item["entity_title"] == "АЛФ")
    assert gap["day"] == "2026-10-30" and gap["deficit"] == "270000.00"
    assert any(hint["kind"] == "transfer" and hint["from_title"] == "БЕТ" for hint in gap["hints"]), "у Беты запас - перевести"
    assert any(hint["kind"] == "move" for hint in gap["hints"]), "платёж до разрыва - можно сдвинуть"

    # Неснижаемый остаток - разрыв глубже на него.
    _one(owner, "cash.settings", reserve={made["alpha"]: "20000"})
    gap = next(item for item in _calendar(owner)["gaps"] if item["entity_title"] == "АЛФ")
    assert gap["deficit"] == "290000.00"

    # Сценарий «только расходы» и строки вне остатка.
    assert _closing(_calendar(owner, scenario="cash"), "АЛФ", "2026-10-30") == "-370000.00"
    assert _closing(_calendar(owner, exclude=["in:debt:"]), "АЛФ", "2026-10-30") == "-370000.00"

    # «Отменить» обещание - долг снова без даты.
    _ok(owner.post(f"{MODEL}/undo", json={"point": promise["point"]["id"]}))
    calendar = _calendar(owner)
    assert _value(calendar, "АЛФ", "in:debt:договоры", "2026-10-29") in (None, "0.00")
    described = _ok(owner.get(f"{MODEL}/describe"))
    assert described["calendar"]["open"] and {"cash.plan", "cash.promise", "cash.import"} <= {item["name"] for item in described["commands"]}
    route = _ok(owner.get(f"{BASE}/core/calendar", params={"from": "2026-10-26", "days": 7}))
    assert route["days"][0] == "2026-10-26" and len(route["days"]) == 7


def _workbook() -> bytes:
    book = openpyxl.Workbook()
    plan = book.active
    plan.title = "План платежей"
    plan.append(["Платежи на ноябрь"])
    plan.append([])
    plan.append(["Дата оплаты", "Фирма", "Строка календаря", "Статья", "Получатель", "Сумма", "Комментарий"])
    plan.append(["02.11.2026", "АЛФ", "Аренда офисов", "Аренда офиса", "ТОО Арендодатель", 450000, ""])
    plan.append(["16.11.2026", "АЛФ", "Зарплата", "Аванс", "Сотрудники", 300000, "средний аванс"])
    plan.append(["17.11.2026", "Неизвестное ТОО", "Прочее", "", "", 1000, ""])
    extra = book.create_sheet("Внеплановые")
    extra.append(["Дата", "Фирма (код)", "Вид", "Сумма", "Что / от кого"])
    extra.append(["28.10.2026", "АЛФ", "Поступление", 1000000, "Займ учредителя"])
    debts = book.create_sheet("Долги клиентов")
    debts.append(["Долги и обещания"])
    debts.append(["Наше юрлицо", "Клиент", "Договор", "Итого долг", "Ожидаемая дата оплаты", "Ожидаемая сумма", "Итоги переговоров"])
    debts.append(["Долги до 60 дней"])
    debts.append(["АЛФ", 'ТОО "Клиент"', "К-1", -100000, "05.11.2026", None, "обещал до 5-го"])
    debts.append(["Итого", None, None, -100000])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _edit(raw: bytes, sheet: str, row: int, column: int, value) -> bytes:
    book = openpyxl.load_workbook(io.BytesIO(raw))
    if value is None and column == 0:
        book[sheet].delete_rows(row)
    else:
        book[sheet].cell(row=row, column=column, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_kalendar_iz_tablicy_faylom(owner: TestClient) -> None:
    made = _company(owner)
    raw = _workbook()
    first = _apply(owner, {"command": "cash.import", "params": {"name": "План.xlsx", "data": base64.b64encode(raw).decode()}})
    out = first["commands"][0]["result"]
    assert {sheet["name"]: sheet["form"] for sheet in out["sheets"]} == {
        "План платежей": "cash_payment", "Внеплановые": "cash_payment", "Долги клиентов": "cash_promise",
    }
    assert (out["new"], out["payments"], out["promises"]) == (4, 3, 1)
    assert (out["income"], out["expense"]) == ("1000000.00", "750000.00")
    assert any("не наше юрлицо" in text for text in out["refusals"]), "чужое юрлицо - отказ словами, а не угадка"

    calendar = _calendar(owner)
    assert _value(calendar, "АЛФ", "out:plan:Аренда офисов", "2026-11-02") == "450000.00"
    assert _value(calendar, "АЛФ", "in:plan:Займ учредителя", "2026-10-28") == "1000000.00"
    # Обещание без суммы - долг из колонки «Итого долг» той же строки.
    assert _value(calendar, "АЛФ", "in:debt:договоры", "2026-11-05") == "100000.00"

    again = _one(owner, "cash.import", name="План.xlsx", data=base64.b64encode(raw).decode())
    assert (again["new"], again["changed"], again["same"], again["removed"]) == (0, 0, 4, 0), "повтор листа не двоит"

    changed = _edit(_edit(raw, "План платежей", 4, 6, 460000), "План платежей", 5, 0, None)
    second = _apply(owner, {"command": "cash.import", "params": {"name": "План.xlsx", "data": base64.b64encode(changed).decode()}})
    result = second["commands"][0]["result"]
    assert (result["changed"], result["removed"]) == (1, 1), "сумма поправлена, строки больше нет - ожидание снято"
    calendar = _calendar(owner)
    assert _value(calendar, "АЛФ", "out:plan:Аренда офисов", "2026-11-02") == "460000.00"
    assert _value(calendar, "АЛФ", "out:plan:Зарплата", "2026-11-16") in (None, "0.00")

    _ok(owner.post(f"{MODEL}/undo", json={"point": second["point"]["id"]}))
    calendar = _calendar(owner)
    assert _value(calendar, "АЛФ", "out:plan:Аренда офисов", "2026-11-02") == "450000.00"
    assert _value(calendar, "АЛФ", "out:plan:Зарплата", "2026-11-16") == "300000.00"
    plans = _ok(owner.get(f"{BASE}/core/calendar/plans"))["plans"]
    assert {plan["source"] for plan in plans} == {"table"} and len(plans) == 4
    assert made["alpha"] in {plan["entity_id"] for plan in plans}


def test_povtor_zhurnala_predlozhenie_i_prinyat(owner: TestClient) -> None:
    made = _company(owner)
    for day, amount in (("2026-07-06", "45000"), ("2026-08-05", "45000"), ("2026-09-07", "47000")):
        _one(owner, "payment.record", direction="out", day=day, amount=amount, entity=made["alpha"], party="ТОО Связь", key=f"net-{day}")
    calendar = _calendar(owner)
    found = [item for item in calendar["proposals"] if item["party"] == "ТОО Связь"]
    assert [(item["day"], item["amount"]) for item in found] == [("2026-11-06", "47000.00")], (
        "3 из 3 месяцев около 6-го: октябрьский срок прошёл, ноябрь - сумма последнего месяца"
    )
    assert "3 из 3" in found[0]["words"]
    habit = _calendar(owner, scenario="habit")
    assert any(item["key"].startswith("out:habit:") for item in habit["lines"]), "по привычке повтор считается сам"
    assert not any(item["key"].startswith("out:habit:") for item in calendar["lines"]), "по условиям - только предложение"

    _one(owner, "cash.accept", keys=[found[0]["key"]])
    calendar = _calendar(owner)
    assert not [item for item in calendar["proposals"] if item["party"] == "ТОО Связь"], "принятое больше не предлагается"
    plan = next(item for item in calendar["items"] if item["source"] == "plan" and item["client"] == "ТОО Связь")
    assert (plan["day"], plan["amount"]) == ("2026-11-06", "47000.00")


def test_srok_stroki_tablicy_i_prazdniki_strany() -> None:
    from app.finance.core import book
    from app.finance.core.cash_calendar import Calendar
    from app.finance.packs import kz

    def row(**values):
        return book.BookRow(line=1, own="АЛФ", customer="Клиент", number="К-1", month=values.pop("month", "2026-11"), month_label="", **values)

    assert Calendar._table_base(row(bill_day="5")) == date(2026, 11, 5)
    assert Calendar._table_base(row(bill_day="31")) == date(2026, 11, 30), "31-е в ноябре - последний день месяца"
    assert Calendar._table_base(row(bill_day="12.11.2026")) == date(2026, 11, 12)
    assert Calendar._table_base(row(bill_day="-", period_from=date(2026, 10, 25))) == date(2026, 11, 25)
    assert Calendar._table_base(row(month="")) is None, "долг до начала учёта - без срока"
    off = kz.COUNTRY.days_off([2026])
    assert date(2026, 10, 26) in off and date(2026, 12, 16) in off
    assert kz.COUNTRY.workday(date(2026, 10, 24), off) == date(2026, 10, 27)
    assert kz.COUNTRY.workday(date(2026, 11, 7), off) == date(2026, 11, 9)
