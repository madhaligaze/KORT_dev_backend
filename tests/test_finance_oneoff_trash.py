"""«Разовые», сводка оплат и корзина (27.09.2026).

* «Оплачено (сводка)» - как колонка Q книги юротдела «Разовые»: сумма «Сумма
  Факт Поступ.» строк основной сводки того же договора. Колонки сводки - по
  названиям; договор - по номеру и клиенту, без угадывания.
* «Разовые» - вид отдела из шаблона отрасли (с 28.10, этап 2, шаг 6): отборы по
  сроку от даты договора и «Остатки» по сводке.
* Корзина: удалённое видно, возвращается, а насовсем не стирается то, на что
  ещё ссылаются живые договоры.
"""
from __future__ import annotations

import uuid
from datetime import timedelta
from decimal import Decimal

import pytest

from app.finance import service as finance_service
from app.finance import trash
from app.finance.contracts import export as export_module
from app.finance.contracts import service, setup, summary, views
from app.finance.contracts.models import Contract
from app.finance.core import book as core_book
from app.finance.db import finance_session

import company_layer


@pytest.fixture(autouse=True)
def _company_layer(monkeypatch):
    """Компания теста - с шапкой таблиц копии прода: её знает словарь компании, а не ядро (п.9 «Пакеты»)."""
    company_layer.everywhere(monkeypatch)


FULL = service.Access(view=True, edit=True, setup=True)
OWNER = service.Actor(uuid.uuid4(), "owner@test")

#: Шапка как у «Сводки все ЮР лица» - с лишними колонками между нужными.
HEADER = [
    "Техн. 2Наша \nФирма", "Техн. 2", "Мес", "Техн. 1", ".", "Наша \nФирма", "Вид \nУслуги", "Наш \nСотрудник",
    "Заказчик\n(Название Фирмы)", "По Выписке\nБанка", "По 1C", "Число", "Сумма\nДоговора", "Отдел", "Предмет",
    "Дата", "№\nДоговора", "Доп", "c", "по", "Коммент", ".", "Сальдо", ".", "Счет", "Счет за", "№ Счета", "Дата выст.",
    "Факт Оплата", "Сумма Факт\nПоступ.",
]


def _row(customer: str, number: str, paid: str, own: str = "BBC", month: str = "АВГУСТ 2026") -> list[str]:
    row = [""] * len(HEADER)
    row[3], row[5], row[8], row[13], row[16], row[29] = month, own, customer, "ЮО", number, paid
    return row


def _put(workspace_id: uuid.UUID, values: list[list[str]]) -> summary.Index:
    """Сводка «прочитана» (шапка копии прода): общий разбор листа лежит у компании, «Оплачено (сводка)» строится по нему."""
    book = core_book.parse(values, title="Сводка", worksheet="Сводка все ЮР лица", columns=core_book.columns_for(company_layer.COLUMNS))
    core_book.put(workspace_id, book)
    return summary.index_of(book)


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    yield engine


@pytest.fixture
def space(finance_db):
    with finance_session() as session:
        workspace = finance_service.ensure_workspace(session)
        setup.add_entity(session, workspace, name="BBC legal support", code="BBCL")
        return workspace.id


def _ws(session, space_id):
    return finance_service.get_workspace(session, space_id)


def _make(session, space_id, **values):
    return service.create(session, _ws(session, space_id), FULL, OWNER, values)


# ── Сводка ──────────────────────────────────────────────────────────────────


def test_svodka_po_nazvaniyam_i_bez_ugadyvaniya():
    index = summary.build_index(
        [
            HEADER,
            _row('Бухгалтерская фирма «Ajour»', "№ЮО/135", "188 754"),
            _row("ТОО Альфа", "№ЮО/1", "100 000", month="ИЮЛЬ 2026"),
            _row("ТОО Альфа", "№ЮО/1", "50 000", month="АВГУСТ 2026"),
            _row("ТОО Izmir Trade", "№ОБО/65", "70 000"),
            _row("ТОО Бета Строй", "№7", "1"),
            _row("ТОО Бета Сервис", "№7", "2"),
        ],
        title="Сводка",
        worksheet="Сводка все ЮР лица",
    )
    # Написание клиента в реестре другое - клиент тот же.
    found = index.match("№ ЮО/135", ["ТОО Бухгалтерская Фирма Ajour"])
    assert (found.state, found.paid) == ("found", Decimal("188754"))
    # Договор растянут по месяцам - оплаты складываются, как в каждой строке.
    spread = index.match("№ЮО/1", ["ТОО Альфа"])
    assert (spread.paid, spread.rows, spread.months) == (Decimal("150000"), 2, ("ИЮЛЬ 2026", "АВГУСТ 2026"))
    # Номер есть, клиент другой - не наш договор, в «Оплачено» ничего.
    other = index.match("№ОБО/65", ["Prosperity KZ audit"])
    assert (other.state, other.paid) == ("other_client", None)
    assert index.match("№404", ["ТОО Альфа"]).state == "missing"
    # Два похожих клиента под одним номером - спорно, не выбираем.
    assert index.match("№7", ["ТОО Бета"]).state == "ambiguous"


def test_svodka_perezhivaet_vstavlennuyu_kolonku():
    shifted = HEADER[:10] + ["Новая колонка"] + HEADER[10:]
    row = _row("ТОО Альфа", "№1", "500")
    row = row[:10] + [""] + row[10:]
    index = summary.build_index([shifted, row], title="", worksheet="Сводка")
    # Колонка находится по названию; позиций у ядра нет (п.9 «Пакеты»), и
    # строки «колонка переехала» больше нет.
    assert index.match("№1", ["ТОО Альфа"]).paid == Decimal("500")


def test_svodka_bez_nuzhnoy_kolonki_otkazyvaet():
    broken = [name for name in HEADER if "Факт" not in name]
    with pytest.raises(summary.SummaryError, match="Сумма Факт"):
        summary.build_index([broken], title="", worksheet="Сводка")


def test_dengi_iz_knigi():
    assert summary.read_money("464 000") == Decimal("464000")
    assert summary.read_money("1 234,50") == Decimal("1234.50")
    assert summary.read_money("-") == Decimal("0")
    assert summary.read_money("(500)") == Decimal("-500")


def test_ssylka_na_knigu():
    assert summary.spreadsheet_id_of(
        "https://docs.google.com/spreadsheets/d/1xEp_QEirE49g/edit?gid=2084632423#gid=2084632423"
    ) == "1xEp_QEirE49g"
    assert summary.spreadsheet_id_of("1xEp_QEirE49g") == "1xEp_QEirE49g"


# ── Числовые условия листа ──────────────────────────────────────────────────


def test_chislovye_usloviya():
    rule = views.validate({"any": [{"all": [{"field": "age_months", "op": "lt", "value": 2}]}]})
    assert views.matches(rule, {"age_months": 1})
    assert not views.matches(rule, {"age_months": 2})
    # Пустое не меньше и не больше ничего: договор без даты - ни в каком сроке.
    assert not views.matches(rule, {})
    with pytest.raises(views.FilterError):
        views.validate({"any": [{"all": [{"field": "age_months", "op": "gte", "value": "много"}]}]})


# ── «Разовые» - вид отдела из шаблона отрасли (этап 2, шаг 6) ───────────────


def _oneoff(session, workspace, code: str = "ЮО") -> tuple[dict, str]:
    """«Разовые договоры по возрасту долга», сохранённые отделом: вид и ключ его книги.

    До шага 6 «Разовые» засевались каждой компании; теперь их сохраняет отдел.
    """
    from app.finance.contracts.models import Department
    from app.finance.core import views as core_views

    registry = service.Registry(session, workspace)
    department = next((item for item in registry.departments.values() if item.code == code), None)
    if department is None:
        department = setup.upsert_department(session, workspace, {"code": code})
    assert isinstance(department, Department)
    viewer = core_views.Viewer(user_id=None, admin=True)
    view = core_views.create(
        session, workspace, viewer, "owner@test", {"template": "oneoff_age", "owner": {"department": str(department.id)}}
    )
    return view, view["key"]


def test_novaya_kompaniya_bez_razovyh(space):
    """Новая компания начинает без «Разовых»: реестр и «Долги», листов чужой книги нет."""
    from app.finance.core import views as core_views

    with finance_session() as session:
        workspace = _ws(session, space)
        registry = service.Registry(session, workspace)
        assert {view.book or "" for view in registry.views} == {""}
        keys = [(view.key, view.title, view.source) for view in core_views.all_of(session, workspace.id)]
        assert keys == [("", "Реестр", "contracts"), ("dolgi", "Долги", "debts")]


def test_razovye_iz_shablona_i_otbirayut(space, monkeypatch):
    with finance_session() as session:
        workspace = _ws(session, space)
        today = service.today()
        common = {"executor": "BBC legal support", "type": "Разовая услуга", "department": "ЮО"}
        view, book = _oneoff(session, workspace)
        assert view["title"] == "Разовые ЮО"
        fresh = _make(session, space, customer="ТОО Альфа", number="№ЮО/1", status="На исполнении", amount="150000",
                      signed_at=(today - timedelta(days=20)).isoformat(), **common)
        old = _make(session, space, customer="ТОО Бета", number="№ЮО/2", status="На исполнении", amount="300000",
                    signed_at=(today - timedelta(days=200)).isoformat(), **common)
        done = _make(session, space, customer="ТОО Гамма", number="№ЮО/3", status="Исполнен", amount="100000",
                     signed_at=(today - timedelta(days=100)).isoformat(), **common)
        _make(session, space, executor="BBC legal support", customer="ТОО Дельта", number="№5",
              type="Абонентское обслуживание", department="ЮО")
        # Разовый договор другого отдела - не в виде ЮО.
        _make(session, space, customer="ТОО Эпсилон", number="№НО/1", status="На исполнении",
              executor="BBC legal support", type="Разовая услуга", department="НО",
              signed_at=(today - timedelta(days=20)).isoformat())

        registry = service.Registry(session, workspace)
        sheets = sorted((view for view in registry.views if view.book == book), key=lambda v: v.position)
        assert [view.title for view in sheets] == ["Все", "до 2 мес", "2-3 мес", "3-6 мес", "6+ мес", "Остатки"]

        # Сводка «прочитана»: Альфа оплатила 150 000 из 150 000, Гамма - 40 000 из 100 000.
        index = _put(workspace.id, [
            HEADER, _row("ТОО Альфа", "№ЮО/1", "150000", own="BBCL"), _row("ТОО Гамма", "№ЮО/3", "40000", own="BBCL"),
        ])
        try:
            listed = service.list_all(session, workspace, FULL)
        finally:
            summary.forget(workspace.id)
        places = {item["values"]["number"]: {place["view"] for place in item["views"]} for item in listed["contracts"]}
        assert places["№ЮО/1"] >= {f"{book}_all", f"{book}_2m"}
        assert f"{book}_6p" in places["№ЮО/2"] and f"{book}_2m" not in places["№ЮО/2"]
        # Исполнен - ни в каком сроке, но с остатком - в «Остатках».
        assert f"{book}_3m" not in places["№ЮО/3"] and f"{book}_rest" in places["№ЮО/3"]
        # Оплачено всё - в «Остатках» нет; в сводке нет - тоже нет (не «ноль»).
        assert f"{book}_rest" not in places["№ЮО/1"]
        assert f"{book}_rest" not in places["№ЮО/2"]
        assert not any(view.startswith(book) for view in places["№5"])
        assert not any(view.startswith(book) for view in places["№НО/1"])
        # Клиент сверяет номер сводки, по которой разложены листы, со свежим.
        assert listed["summary_rev"] == index.rev
        assert {fresh.id, old.id, done.id}


def test_razovye_dva_statusa_i_zelyonaya_stroka(space):
    """28.09.2026: в «Разовых» выбор статуса - «на исполнении» и «исполнен», строка
    «исполнен» подсвечена (как `=TRIM($E2)="исполнен"` в книге юротдела)."""
    with finance_session() as session:
        workspace = _ws(session, space)
        _view, book = _oneoff(session, workspace)
        common = {"executor": "BBC legal support", "type": "Разовая услуга", "department": "ЮО"}
        done = _make(session, space, customer="ТОО Гамма", number="№ЮО/3", status="Исполнен", **common)
        going = _make(session, space, customer="ТОО Альфа", number="№ЮО/1", status="На исполнении", **common)
        registry = service.Registry(session, workspace)
        statuses = {value.value: str(value.id) for value in registry.values.values() if value.field_key == "status"}
        main = next(view for view in registry.views if view.key == f"{book}_all")
        block = main.blocks[0]
        assert block["choices"] == {"status": [statuses["На исполнении"], statuses["Исполнен"]]}
        assert block["defaults"]["status"] == statuses["На исполнении"]

        listed = {item["id"]: item for item in service.list_all(session, workspace, FULL)["contracts"]}
        tone = {place["view"]: place.get("tone") for place in listed[str(done.id)]["views"]}
        assert tone[f"{book}_all"] == "done"
        assert all("tone" not in place for place in listed[str(going.id)]["views"])
        # В реестре подсветки нет - правило живёт у блока вида.
        assert "tone" not in next(place for place in listed[str(done.id)]["views"] if place["view"] == "main")

        # Выбор - только значения списков и только у полей-списков.
        with pytest.raises(finance_service.FinanceError):
            setup.upsert_view(session, workspace, {"blocks": [{**block, "choices": {"number": [statuses["Исполнен"]]}}]}, main.id)
        with pytest.raises(finance_service.FinanceError):
            setup.upsert_view(session, workspace, {"blocks": [{**block, "choices": {"status": [str(uuid.uuid4())]}}]}, main.id)
        with pytest.raises(views.FilterError):
            setup.upsert_view(session, workspace, {"blocks": [{**block, "paint": [{"filter": block["filter"], "tone": "red"}]}]}, main.id)
        # Подсветка без условий не хранится - она не красила бы ни одной строки.
        cleared = setup.upsert_view(session, workspace, {"blocks": [{**block, "paint": [{"filter": {"any": []}}]}]}, main.id)
        assert "paint" not in cleared.blocks[0]


def test_razovye_tolko_dva_statusa(space):
    """30.09.2026: «Не состоялся» стоял в «Разовых ЮО», хотя книга юротдела держит
    только «на исполнении» и «исполнен». Решение владельца: в реестре - все
    договоры, в «Разовых» - только эти два статуса, без статуса тоже нет."""
    with finance_session() as session:
        workspace = _ws(session, space)
        _view, book = _oneoff(session, workspace)
        setup.add_value(session, workspace, "status", "нужно закрыть по бух")
        signed = (service.today() - timedelta(days=20)).isoformat()
        for number, status in (
            ("№ЮО/1", "На исполнении"), ("№ЮО/2", "Исполнен"), ("№ЮО/3", "Не состоялся"),
            ("№ЮО/4", "Недействующий"), ("№ЮО/5", "нужно закрыть по бух"), ("№ЮО/6", ""),
        ):
            _make(session, space, executor="BBC legal support", customer=f"ТОО {number}", number=number,
                  type="Разовая услуга", status=status, signed_at=signed, department="ЮО")
        listed = service.list_all(session, workspace, FULL)["contracts"]
        places = {item["values"]["number"]: {place["view"] for place in item["views"]} for item in listed}
        oneoff = {number: {view for view in views if view.startswith(book)} for number, views in places.items()}
        assert oneoff["№ЮО/1"] == {f"{book}_all", f"{book}_2m"}
        assert oneoff["№ЮО/2"] == {f"{book}_all"}
        for number in ("№ЮО/3", "№ЮО/4", "№ЮО/5", "№ЮО/6"):
            assert oneoff[number] == set(), number
        # Реестр держит всё.
        assert all("main" in views for views in places.values())


def test_razovye_otdela_summa_i_ostatok_dolej(space):
    """30.09.2026: в «Разовых ЮО» №ОКР/59 на 348 000 (HR 70% / ЮО 30%) стоял
    суммой 348 000. Юротделу нужна своя доля - 104 400, «Оплачено» и «Остаток» -
    той же долей, а в названии части «Остатков» - итог, как в книге юротдела:
    «· договоров: N · остаток: … ₸». С шага 6 отдел доли - владелец вида."""
    import io

    from openpyxl import load_workbook

    from app.finance.contracts import shares
    from app.finance.core import views as core_views

    admin = service.Access(view=True, edit=True, setup=True, admin=True)
    with finance_session() as session:
        workspace = _ws(session, space)
        yuo = setup.upsert_department(session, workspace, {"code": "ЮО"})
        hr = setup.upsert_department(session, workspace, {"code": "HR"})
        yuo_id = str(yuo.id)
        signed = (service.today() - timedelta(days=20)).isoformat()
        common = {"executor": "BBC legal support", "type": "Разовая услуга", "status": "На исполнении", "signed_at": signed}
        split = _make(session, space, customer="ТОО Акжол", number="№ОКР/59", amount="348000", department="ЮО", **common)
        _make(session, space, customer="ТОО Альфа", number="№ЮО/1", amount="150000", department="ЮО", **common)
        _make(session, space, customer="ТОО Бета", number="№ЮО/2", amount="90000", department="ЮО, HR", **common)
        shares.set_departments(
            session, workspace, admin, OWNER, split.id, "percent",
            [{"department_id": str(hr.id), "value": "70"}, {"department_id": str(yuo.id), "value": "30"}],
        )
        # Вида отдела нет - доли нет ни у какой книги.
        assert service.Registry(session, workspace).share_departments() == []
        view, book = _oneoff(session, workspace)
        registry = service.Registry(session, workspace)
        assert registry.share_departments() == [yuo.id]
        keys = [sheet.key for sheet in registry.views if sheet.book == book]

        # Сводка: Акжол заплатил 174 000 из 348 000, Альфа - 50 000 из 150 000.
        _put(workspace.id, [
            HEADER,
            _row("ТОО Акжол", "№ОКР/59", "174000", own="BBCL"),
            _row("ТОО Альфа", "№ЮО/1", "50000", own="BBCL"),
        ])
        try:
            listed = {item["values"]["number"]: item for item in service.list_all(session, workspace, FULL)["contracts"]}
            data = export_module.build(session, workspace, FULL, OWNER, keys)
            # «Спросить» вид - те же числа, что в листе и выгрузке.
            asked = core_views.ask(session, workspace, admin, core_views.find(session, workspace, book))
        finally:
            summary.forget(workspace.id)

    # В ответе реестра - доля ЮО, и только её; договор одного отдела и отделы без долей - без неё.
    assert listed["№ОКР/59"]["department_share"] == {yuo_id: {"amount": "104400.00", "percent": "30"}}
    assert "department_share" not in listed["№ЮО/1"] and "department_share" not in listed["№ЮО/2"]

    book_file = load_workbook(io.BytesIO(data))
    rows = [[cell.value for cell in row] for row in book_file["Остатки"].iter_rows()]
    # Итог части - по долям: 52 200 у ОКР/59 (30% от 174 000) и 100 000 у Альфы.
    assert rows[0][1] == "Работа идёт - есть остаток · договоров: 2 · остаток: 152 200 ₸"
    head = rows[1]
    number, amount, paid, rest = (head.index(label) for label in ("№ Договора", "Сумма Договора", "Оплачено", "Остаток"))
    values = {row[number]: (row[amount], row[paid], row[rest]) for row in rows[2:] if row and row[number]}
    assert values["№ОКР/59"] == (104400, 52200, 52200)
    assert values["№ЮО/1"] == (150000, 50000, 100000)
    # Во всём списке вида: отделов два, доли ЮО нет - договор целиком.
    main = [[cell.value for cell in row] for row in book_file["Все"].iter_rows()]
    column = main[0].index("Сумма Договора")
    assert {row[main[0].index("№ Договора")]: row[column] for row in main[1:]}["№ЮО/2"] == 90000

    sheet = {item["key"]: item for item in asked["sheets"]}
    assert sheet[f"{book}_rest"]["blocks"][0] == {
        "title": "Работа идёт - есть остаток", "contracts": 2, "remaining": "152200",
    }
    measured = {row["number"]: row["measures"] for row in asked["rows"]}
    assert measured["№ОКР/59"]["amount"] == "104400" and measured["№ОКР/59"]["summary_remaining"] == "52200"
    staff = asked["groups"]["people"]
    assert staff["totals"]["contracts"] == 3 and staff["totals"]["amount"] == "344400"


def test_vygruzka_bez_vybora_tolko_reestr(space):
    with finance_session() as session:
        workspace = _ws(session, space)
        _oneoff(session, workspace)
        _make(session, space, executor="BBC legal support", customer="ТОО Альфа", type="Разовая услуга", department="ЮО")
        from openpyxl import load_workbook
        import io

        book = load_workbook(io.BytesIO(export_module.build(session, workspace, FULL, OWNER)))
        assert not any(title in book.sheetnames for title in ("Все", "Остатки"))


# ── Корзина ─────────────────────────────────────────────────────────────────


def test_korzina_dogovor_tuda_i_obratno(space):
    with finance_session() as session:
        workspace = _ws(session, space)
        contract = _make(session, space, executor="BBC legal support", customer="ТОО Альфа", number="№9")
        service.remove(session, workspace, FULL, OWNER, contract.id)
        items = trash.listing(session, workspace)["items"]
        assert [(item["kind"], item["title"]) for item in items] == [("contract", "№9 · ТОО Альфа")]
        seq_before = session.get(Contract, contract.id).seq
        trash.restore(session, workspace, "contract", contract.id)
        restored = session.get(Contract, contract.id)
        assert restored.deleted_at is None and restored.seq > seq_before
        assert trash.listing(session, workspace)["items"] == []


def test_korzina_ne_stiraet_ispolzuemoe(space):
    with finance_session() as session:
        workspace = _ws(session, space)
        _make(session, space, executor="BBC legal support", customer="ТОО Альфа", status="На исполнении")
        registry = service.Registry(session, workspace)
        used = registry.values_by_field["status"]["на исполнении"]
        # Стоящий в договоре статус и в корзину не уходит - сводят с другим.
        with pytest.raises(Exception, match="сведите"):
            setup.update_value(session, workspace, used.id, {"archived": True})
        spare = setup.add_value(session, workspace, "status", "Лишний")
        setup.update_value(session, workspace, spare.id, {"archived": True})
        assert [item["title"] for item in trash.listing(session, workspace)["items"]] == ["Текущее состояние: Лишний"]
        trash.purge(session, workspace, "value", spare.id)
        assert trash.listing(session, workspace)["items"] == []

        # Отдел с договором насовсем не стирается - стёрся бы у договора.
        dept = setup.upsert_department(session, workspace, {"code": "ЮО"})
        _make(session, space, executor="BBC legal support", customer="ТОО Бета", department="ЮО")
        dept.archived_at = service._now()
        session.flush()
        with pytest.raises(trash.TrashError, match="договоров 1"):
            trash.purge(session, workspace, "department", dept.id)


def test_korzina_svoe_pole_nasovsem_so_znacheniyami(space):
    with finance_session() as session:
        workspace = _ws(session, space)
        field = setup.add_field(session, workspace, title="Источник", type="text")
        contract = _make(session, space, executor="BBC legal support", customer="ТОО Альфа", **{field.key: "ОМиП"})
        assert (session.get(Contract, contract.id).attrs or {}).get(field.key) == "ОМиП"
        setup.update_field(session, workspace, field.key, {"archived": True})
        trash.purge(session, workspace, "field", field.id)
        assert field.key not in (session.get(Contract, contract.id).attrs or {})


def test_korzina_sistemnoe_ne_udalyaetsya(space):
    with finance_session() as session:
        workspace = _ws(session, space)
        registry = service.Registry(session, workspace)
        main = next(view for view in registry.views if view.main)
        with pytest.raises(trash.TrashError):
            trash.restore(session, workspace, "view", main.id)
        with pytest.raises(trash.TrashError):
            trash.purge(session, workspace, "nothing", uuid.uuid4())


# ── Личный вид листа ────────────────────────────────────────────────────────


def test_lichnyy_vid_u_kazhdogo_svoy(finance_db):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.api.routes.finance import router as finance_router
    from app.api.routes.finance_looks import router as looks_router
    from app.api.routes.finance_people import router as people_router

    app = FastAPI()
    for router in (finance_router, people_router, looks_router):
        app.include_router(router, prefix="/api/v1")
    base = "/api/v1/finance"
    owner = TestClient(app)
    registered = owner.post(f"{base}/auth/register", json={
        "company": "Вид", "full_name": "Владелец Вида", "email": "look@test.local", "password": "pass-12345",
    })
    assert registered.status_code in (200, 201), registered.text
    look = {"v": 1, "sheets": {"main": {"cols": {"customer": {"w": 320}}, "cells": {"c:1|amount": {"bl": 1}}}}}
    assert owner.put(f"{base}/looks/registry", json={"look": look}).json() == {"ok": True}
    assert owner.get(f"{base}/looks/registry").json()["look"] == look
    assert owner.get(f"{base}/looks/journal").json()["look"] == {}
    assert owner.get(f"{base}/looks/..bad").status_code == 404
    assert owner.put(f"{base}/looks/registry", json={"look": {"x": "я" * 600_000}}).status_code == 413
    assert owner.delete(f"{base}/looks/registry").json() == {"ok": True}
    assert owner.get(f"{base}/looks/registry").json()["look"] == {}
