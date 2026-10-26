"""Полигон сам себе не врёт: правда сходится, номера годные, каждый случай каталога есть в данных.

Это проверка генератора, а не KORT. KORT на полигоне проверяют наборы этапа 2.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl
from polygon import sources
from polygon.cases import ASK, CASES, OFFER, REFUSE, SAM
from polygon.world import SINCE, TODAY, build, case_index

from app.finance.core import bin as bin_module


def test_mir_postroen_odinakovo() -> None:
    first, second = build(), build()
    assert sources.truth(first) == sources.truth(second)
    assert sources.truth(build(7)) != sources.truth(first)


def test_kazhdyi_sluchai_kataloga_est_v_dannyh() -> None:
    world = build()
    where = case_index(world)
    # О12 и Д20 - свойства всех данных юрлица, а не одной записи; Ф-случаи живут в файлах.
    in_files = {"О12", "Д19", "Д20", "Ф1", "Ф3", "Ф4", "Ф5", "Ф6", "Ф7"}
    missing = [case.id for case in CASES if not where.get(case.id) and case.id not in in_files]
    assert missing == []
    assert {case.expect for case in CASES} == {SAM, OFFER, ASK, REFUSE}


def test_nomera_godnye_krome_opechatki() -> None:
    world = build()
    for entity in world.entities.values():
        assert bin_module.check(entity.number, entity.name).ok, entity.name
    for party in world.parties.values():
        assert bin_module.check(party.number, party.name).ok, party.name
        if party.registry_number and party.key != "samal":
            assert party.registry_number == party.number
    samal = world.parties["samal"]
    assert not bin_module.check(samal.registry_number).ok
    assert bin_module.kind_of(world.entities["ИП"].number, world.entities["ИП"].name) == bin_module.IP


def test_dolgi_skladyvayutsya() -> None:
    world = build()
    # Предоплата за год: пять месяцев начислено, остальное - аванс.
    assert world.debt("zhuldyz") == Decimal(5 * 120000 - 1440000)
    # Каспий Трейд: апрель и май не оплачены на 01.06, апрель погашен в июле.
    assert world.opening("kaspiy") == Decimal(500000) + _utilities(world, "kaspiy", (4, 5))
    # Летняя пауза: июля и августа нет.
    months = {a.month.month for a in world.accruals("miras") if a.month.year == 2026}
    assert 7 not in months and 8 not in months and 9 in months
    # Перевод на «Тумар Плюс»: октябрь начисляет Плюс.
    assert [a.executor for a in world.accruals("teniz") if a.month == date(2026, 10, 1)] == ["ТП"]
    # Возврат переплаты возвращает долг к нулю за июль.
    assert world.debt("nurtau", date(2026, 7, 31)) == Decimal(-50000)
    assert world.debt("nurtau", date(2026, 8, 15)) - world.debt("nurtau", date(2026, 8, 14)) == Decimal(50000)
    owed = world.owed()
    assert all(owed[code] > 0 for code in ("ТС", "ТО", "ИП")) and all(value >= 0 for value in owed.values())
    assert SINCE < TODAY


def _utilities(world, key: str, months: tuple[int, ...]) -> Decimal:
    return sum((a.utilities for a in world.accruals(key) if a.month.year == 2026 and a.month.month in months), Decimal(0))


def test_istochniki_otkryvayutsya() -> None:
    world = build()
    files = sources.render(world)
    for name, content in files.items():
        if name.endswith(".xlsx"):
            openpyxl.load_workbook(io.BytesIO(content))
    rent = openpyxl.load_workbook(io.BytesIO(files["Excel/Аренда - начисления 2026.xlsx"]))
    assert "Сентябрь" not in rent.sheetnames and "Октябрь" in rent.sheetnames
    october = [cell.value for cell in rent["Октябрь"][3]]
    assert october.count("Сумма") == 2
    registry = openpyxl.load_workbook(io.BytesIO(files["Excel/Реестр договоров 2026.xlsx"]))
    hidden = [row for row, dim in registry["Клининг"].row_dimensions.items() if dim.hidden]
    assert len(hidden) == 1
    odata = files["1С Тумар Сервис/odata.json"]
    taza = [row for row in odata["Catalog_Контрагенты"] if row["ИдентификационныйКодЛичности"] == world.parties["taza"].number]
    assert len(taza) == 2


def test_snimok_1s_chitaet_podklyuchenie() -> None:
    """Снимок OData читается тем же `onec/sync.read`, что настоящая база."""
    from app.finance.onec import sync as onec_sync

    class Fake:
        def __init__(self, data):
            self.data = data

        def rows(self, entity, **_):
            yield from self.data.get(entity, [])

    world = build()
    odata = sources.onec_odata(world, "ТС")
    org = odata["Catalog_Организации"][0]["Ref_Key"]
    snapshot = onec_sync.read(Fake(odata), organization_ref=org, since=date(2026, 1, 1))
    counts = snapshot.counts()
    assert counts["sale"] == sum(1 for d in world.documents if d.entity == "ТС" and d.kind == "act")
    assert counts["invoice"] == sum(1 for d in world.documents if d.entity == "ТС" and d.kind == "invoice")
    assert counts["payment"] >= sum(1 for p in world.payments if p.dst == "ТС" and p.kind == "contract" and p.day.year == 2026)
    unsigned = [d for d in snapshot.documents if d["kind"] == "sale" and d["party_bin"] == world.parties["samal"].number and not d.get("signed_at")]
    assert len(unsigned) >= 3


def test_vypiska_kaspi_razbiraetsya() -> None:
    """Выписки полигона - в форме Kaspi Business, которую разбор KORT уже знает."""
    from app.finance.importing import analyze

    world = build()
    data, _closing = sources.kaspi_statement(world, "ТС", date(2026, 8, 1), Decimal(1000000))
    entity = world.entities["ТС"]
    preview = analyze(data, "Выписка_по_счету.xlsx", ["Kaspi Тумар Сервис"], account_numbers={entity.iban: "Kaspi Тумар Сервис"})
    expected = sum(1 for p in world.payments if date(2026, 8, 1) <= p.day <= date(2026, 8, 31) and "ТС" in (p.src, p.dst))
    assert preview.question is None
    taken = [row for row in preview.rows if row.state != "skipped"]
    assert len(taken) == expected, [(row.line, row.state, row.problems) for row in preview.rows if row.state != "imported"]
