"""Источники полигона: как правду группы «Тумар» видят 1С, Excel отделов и банк.

Каждый источник рисуется так, как его отдаёт живая программа или ведёт живой
человек, вместе с искажениями из каталога (`cases.py`):

* «Тумар Сервис» - 1С в интернете: снимок для `FakeClient` в форме OData
  (те же сущности и поля, что читает `onec/sync.py`) и ОСВ по 1210 файлом;
* «Тумар Офис» и «Тумар Плюс» - одна 1С на компьютере бухгалтера: только
  файлы (контрагенты, журнал реализаций, журнал счетов, ОСВ по 1210);
* реестр договоров офис-менеджера, таблица начислений аренды, таблица
  слушателей учебного центра - Excel отделов;
* выписки Kaspi Business в Excel по каждому счёту и месяцу.

`render(world)` отдаёт словарь «путь → содержимое» (bytes или dict для JSON).
"""
from __future__ import annotations

import io
import json
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font

from polygon.cases import CASES
from polygon.world import (
    CLEANING,
    RENT,
    SCHOOL,
    SERVICE,
    SINCE,
    TODAY,
    Contract,
    World,
    case_index,
    month_end,
    month_range,
)

SALE_TYPE = "StandardODATA.Document_РеализацияТоваровУслуг"
SF_TYPE = "StandardODATA.Document_СчетФактураВыданный"
_NS = uuid.UUID("6f0c9d1e-26a0-4c3b-9e1a-261000000000")
BIC = {"722": "CASPKZKA", "601": "HSBKKZKX", "856": "KCJBKZKX"}
MONTH_SHEETS = {6: "Июнь", 7: "Июль", 8: "Август", 9: "Сентябрь", 10: "Октябрь"}


def ref(*parts: Any) -> str:
    """Постоянный Ref_Key 1С для вещи полигона: тот же мир - те же ключи."""
    return str(uuid.uuid5(_NS, "/".join(str(part) for part in parts)))


def iso(day: date, hour: int = 9) -> str:
    return f"{day:%Y-%m-%d}T{hour:02d}:00:00"


def _xlsx(wb: openpyxl.Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _onec_contract(contract: Contract) -> tuple[str, str]:
    """Договор в 1С: (ключ, подпись). Не заведён - документы идут на «Основной договор» клиента."""
    if contract.onec_number == "":
        return ref("contract-main", contract.party), "Основной договор"
    number = contract.onec_number or contract.number
    return ref("contract", contract.key), f"Договор № {number} от {contract.signed:%d.%m.%Y}"


def _party_label(world: World, key: str) -> tuple[str, str, str]:
    """Сторона в 1С: (имя, полное имя, БИН/ИИН); у нашего юрлица - его имя."""
    if key in world.entities:
        entity = world.entities[key]
        return entity.name, entity.name, entity.number
    party = world.parties[key]
    full = party.renamed[1] if party.renamed else party.name
    return party.onec_name, full, party.number


# ── 1С в интернете: снимок OData ─────────────────────────────────────────────


def onec_odata(world: World, entity_key: str = "ТС") -> dict[str, list[dict[str, Any]]]:
    """База 1С юрлица в форме REST-сервиса: сущность → строки, как их читает `onec/sync.py`."""
    entity = world.entities[entity_key]
    org = ref("org", entity_key)
    docs = [doc for doc in world.documents if doc.entity == entity_key]
    contracts = {doc.contract for doc in docs}
    payments = [p for p in world.payments if p.dst == entity_key and p.kind == "contract" and p.day >= date(2026, 1, 1)]
    party_keys = {world.contracts[c].party for c in contracts} | {p.src for p in payments}
    party_keys -= set(world.entities)

    parties = []
    for key in sorted(party_keys):
        party = world.parties[key]
        parties.append({
            "Ref_Key": ref("party", key), "Description": party.onec_name, "НаименованиеПолное": _party_label(world, key)[1],
            "ИдентификационныйКодЛичности": party.number, "IsFolder": False, "DeletionMark": False,
        })
        if party.onec_duplicate:
            parties.append({
                "Ref_Key": ref("party-dup", key), "Description": party.onec_duplicate, "НаименованиеПолное": party.name,
                "ИдентификационныйКодЛичности": party.number, "IsFolder": False, "DeletionMark": False,
            })

    def party_ref(key: str, day: date) -> str:
        party = world.parties[key]
        return ref("party-dup", key) if party.onec_duplicate and day < date(2026, 5, 1) else ref("party", key)

    catalog = []
    for key in sorted(party_keys):
        catalog.append({"Ref_Key": ref("contract-main", key), "Owner_Key": ref("party", key), "НомерДоговора": "",
                        "ДатаДоговора": "0001-01-01T00:00:00", "Description": "Основной договор", "IsFolder": False})
    for key in sorted(contracts):
        contract = world.contracts[key]
        if contract.onec_number == "":
            continue
        contract_ref, label = _onec_contract(contract)
        catalog.append({"Ref_Key": contract_ref, "Owner_Key": ref("party", contract.party), "НомерДоговора": contract.onec_number or contract.number,
                        "ДатаДоговора": iso(contract.signed, 0), "Description": label, "IsFolder": False})

    def head(doc, kind: str, hour: int) -> dict[str, Any]:
        contract = world.contracts[doc.contract]
        return {
            "Ref_Key": ref(kind, doc.key), "DataVersion": "AAAAAQAAAAA=", "Number": doc.number, "Date": iso(doc.day, hour),
            "Posted": True, "DeletionMark": False, "Организация_Key": org, "Контрагент_Key": party_ref(contract.party, doc.day),
            "ДоговорКонтрагента_Key": _onec_contract(contract)[0], "СуммаДокумента": float(doc.amount_1c),
        }

    invoices, sales, signed, sfs, esfs = [], [], [], [], []
    for doc in docs:
        if doc.kind == "invoice":
            invoices.append({**head(doc, "invoice", 9), "Услуги": [{"LineNumber": "1", "Содержание": doc.content}]})
            continue
        row = {**head(doc, "sale", 18), "ВидОперации": "Услуги", "Услуги": [{"LineNumber": "1", "Содержание": doc.content}]}
        if world.contracts[doc.contract].billing == "month":
            row["ДатаНачалаОтчетногоПериода"] = iso(doc.month, 0)
            row["ДатаОкончанияОтчетногоПериода"] = iso(month_end(doc.month), 0)
        sales.append(row)
        if doc.signed:
            signed.append({"Организация_Key": org, "ДокументРеализацииУслуг": ref("sale", doc.key),
                           "ДокументРеализацииУслуг_Type": SALE_TYPE, "ДатаПодписания": iso(doc.signed, 0)})
        if entity.vat_from and doc.day >= entity.vat_from:
            sf = {**head(doc, "sf", 19), "ДокументОснование": ref("sale", doc.key), "ДокументОснование_Type": SALE_TYPE,
                  "СпособВыставления": "Электронно"}
            sfs.append(sf)
            if doc.esf:
                esfs.append({**head(doc, "esf", 12), "Date": iso(doc.esf, 12), "Направление": "Исходящий", "Статус": "Принят",
                             "Состояние": "Принят", "РегистрационныйНомер": f"KZ-ESF-{entity.number}-{doc.number[-5:]}",
                             "СчетФактура": sf["Ref_Key"], "СчетФактура_Type": SF_TYPE, "ДатаОборота": iso(doc.day, 0)})

    incoming = []
    for payment in payments:
        lines: dict[str, Decimal] = {}
        for contract_key, _month, amount in payment.alloc:
            contract_ref = _onec_contract(world.contracts[contract_key])[0]
            lines[contract_ref] = lines.get(contract_ref, Decimal(0)) + amount
        incoming.append({
            "Ref_Key": ref("payment", payment.key), "DataVersion": "AAAAAQAAAAA=", "Number": f"{len(incoming) + 1:011d}",
            "Date": iso(payment.day, 12), "Posted": True, "DeletionMark": False, "Организация_Key": org, "ВидОперации": "ОплатаПокупателя",
            "Контрагент_Key": party_ref(payment.src, payment.day), "СуммаДокумента": float(payment.amount), "НазначениеПлатежа": payment.purpose,
            "РасшифровкаПлатежа": [
                {"LineNumber": str(index + 1), "ДоговорКонтрагента_Key": contract_ref, "СуммаПлатежа": float(amount),
                 "СчетНаОплату_Key": "00000000-0000-0000-0000-000000000000"}
                for index, (contract_ref, amount) in enumerate(lines.items())
            ],
        })

    # Возврат переплаты клиенту - «Платёжное поручение исходящее» (О6): в 1С он
    # лежит рядом с поступлениями, по договору клиента.
    refunds = []
    for payment in (p for p in world.payments if p.src == entity_key and p.kind == "refund"):
        contract_ref = _onec_contract(world.contracts[payment.alloc[0][0]])[0]
        refunds.append({
            "Ref_Key": ref("refund", payment.key), "DataVersion": "AAAAAQAAAAA=", "Number": f"{len(refunds) + 1:011d}",
            "Date": iso(payment.day, 12), "Posted": True, "DeletionMark": False, "Организация_Key": org,
            "ВидОперации": "ВозвратДенежныхСредствПокупателю", "Контрагент_Key": party_ref(payment.dst, payment.day),
            "СуммаДокумента": float(payment.amount), "НазначениеПлатежа": payment.purpose,
            "РасшифровкаПлатежа": [{"LineNumber": "1", "ДоговорКонтрагента_Key": contract_ref, "СуммаПлатежа": float(payment.amount),
                                    "СчетНаОплату_Key": "00000000-0000-0000-0000-000000000000"}],
        })

    return {
        "Catalog_Организации": [{"Ref_Key": org, "Description": entity.name, "ИдентификационныйНомер": entity.number, "DeletionMark": False}],
        "Catalog_Контрагенты": parties,
        "Catalog_ДоговорыКонтрагентов": catalog,
        "Document_СчетНаОплатуПокупателю": invoices,
        "Document_РеализацияТоваровУслуг": sales,
        "InformationRegister_СведенияПоАктамВыполненныхРабот": signed,
        "Document_СчетФактураВыданный": sfs,
        "Document_ЭСФ": esfs,
        "Document_ПлатежноеПоручениеВходящее": incoming,
        "Document_ПлатежноеПоручениеИсходящее": refunds,
    }


# ── 1С файлами: отчёты и списки ──────────────────────────────────────────────


def osv_1210(world: World, entities: tuple[str, ...], month: date = SINCE) -> bytes:
    """«Оборотно-сальдовая ведомость по счету 1210» за месяц по контрагентам и договорам.

    Сальдо - по актам 1С и оплатам; начало месяца - то, что KORT берёт входящим
    сальдо. Покупки сюда не входят (это 3310), между своими - входят: для 1С
    своё ТОО - такой же покупатель.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TDSheet"
    owner = world.entities[entities[0]]
    ws.append([owner.name])
    ws.append([f"Оборотно-сальдовая ведомость по счету 1210 за {MONTH_SHEETS.get(month.month, month.month)} {month.year} г."])
    ws.append(["Выводимые данные: БУ (данные бухгалтерского учета)"])
    ws.append([f"Отбор: Организация В списке \"{'; '.join(world.entities[e].name for e in entities)}\""])
    ws.append([])
    ws.append(["Счет", "Сальдо на начало периода", None, "Обороты за период", None, "Сальдо на конец периода", None])
    ws.append([None, "Дебет", "Кредит", "Дебет", "Кредит", "Дебет", "Кредит"])
    ws.append(["Контрагенты"])
    ws.append(["Договоры"])
    for cells in ("B6:C6", "D6:E6", "F6:G6"):
        ws.merge_cells(cells)

    def split(value: Decimal) -> tuple[Any, Any]:
        return (float(value), None) if value > 0 else (None, float(-value)) if value < 0 else (None, None)

    by_party: dict[str, list[tuple[str, Decimal, Decimal, Decimal]]] = {}
    end = month_end(month)
    for key, contract in world.contracts.items():
        if contract.side == "purchase" or contract.entity not in entities:
            continue
        opening = world.debt(key, month - timedelta(days=1))
        debit = sum((doc.amount_1c for doc in world.documents if doc.contract == key and doc.kind == "act" and month <= doc.day <= end), Decimal(0))
        credit = sum((a for p in world.payments if month <= p.day <= end for c, _m, a in p.alloc if c == key), Decimal(0))
        if not (opening or debit or credit):
            continue
        by_party.setdefault(contract.party, []).append((_onec_contract(contract)[1], opening, debit, credit))
    total = [Decimal(0)] * 3
    rows = []
    for party_key in sorted(by_party, key=lambda k: _party_label(world, k)[0]):
        items = by_party[party_key]
        sums = [sum((item[i] for item in items), Decimal(0)) for i in (1, 2, 3)]
        total = [a + b for a, b in zip(total, sums)]
        rows.append((_party_label(world, party_key)[0], *sums, 1))
        rows.extend((label, opening, debit, credit, 2) for label, opening, debit, credit in items)
    ws.append(["1210", *split(total[0]), float(total[1]) or None, float(total[2]) or None, *split(total[0] + total[1] - total[2])])
    for label, opening, debit, credit, level in rows:
        ws.append([label, *split(opening), float(debit) or None, float(credit) or None, *split(opening + debit - credit)])
        ws.cell(ws.max_row, 1).alignment = Alignment(indent=level)
    ws.append(["Итого", *split(total[0]), float(total[1]) or None, float(total[2]) or None, *split(total[0] + total[1] - total[2])])
    return _xlsx(wb)


def onec_parties(world: World, entities: tuple[str, ...]) -> bytes:
    """Справочник «Контрагенты» 1С, выгруженный списком."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Контрагенты"
    ws.append(["Наименование", "Полное наименование", "ИИН/БИН", "Вид контрагента"])
    keys = {c.party for c in world.contracts.values() if c.entity in entities or any(code in entities for _d, code in c.executors)}
    keys |= {p.src for p in world.payments if p.dst in entities} | {p.dst for p in world.payments if p.src in entities}
    for key in sorted(keys - set(entities), key=lambda k: _party_label(world, k)[0]):
        name, full, number = _party_label(world, key)
        kind = "Юридическое лицо" if (key in world.entities or world.parties[key].kind == "org") else "Физическое лицо"
        ws.append([name, full, number, kind])
    return _xlsx(wb)


def onec_journal(world: World, entities: tuple[str, ...], kind: str) -> bytes:
    """Журнал документов 1С, выгруженный списком: «Реализация товаров и услуг» или «Счета на оплату»."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Реализации" if kind == "act" else "Счета"
    ws.append(["Дата", "Номер", "Организация", "Контрагент", "Договор", "Сумма", "Валюта", "Содержание"])
    for doc in sorted((d for d in world.documents if d.kind == kind and d.entity in entities), key=lambda d: (d.day, d.entity, d.number)):
        contract = world.contracts[doc.contract]
        ws.append([datetime(doc.day.year, doc.day.month, doc.day.day, 18 if kind == "act" else 9), doc.number, world.entities[doc.entity].name,
                   _party_label(world, contract.party)[0], _onec_contract(contract)[1], float(doc.amount_1c), "KZT", doc.content])
        ws.cell(ws.max_row, 1).number_format = "dd.mm.yyyy h:mm:ss"
    return _xlsx(wb)


# ── Excel отделов ────────────────────────────────────────────────────────────


def registry(world: World) -> bytes:
    """Реестр договоров офис-менеджера: лист на отдел, шапка в две строки, итоги, скрытый черновик."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = ((CLEANING, "Клининг", "Сумма в месяц, тг"), (SERVICE, "Техобслуживание", "Сумма в месяц, тг"),
              (RENT, "Аренда", "Аренда в месяц, тг"), (SCHOOL, "Учебный центр", "Стоимость, тг"), ("purchase", "Поставщики", "Сумма в месяц, тг"))
    text_amounts = {"alem": "150 000,00 тг", "orken_media": "150.000"}
    for department, title, money in sheets:
        ws = wb.create_sheet(title)
        ws.append([f"РЕЕСТР ДОГОВОРОВ: {title.upper()}"])
        ws.append([f"на {TODAY - timedelta(days=5):%d.%m.%Y}"])
        ws.append([])
        client = "Поставщик" if department == "purchase" else "Клиент"
        head = ["№", client, "БИН/ИИН", "Договор", None, money, "Наше ТОО", "Помещение" if department == RENT else "Отдел",
                "Ответственный", "Статус", "Примечание"]
        ws.append(head)
        ws.append([None, None, None, "№", "дата", None, None, None, None, None, None])
        for col in (1, 2, 3, 6, 7, 8, 9, 10, 11):
            ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
        ws.merge_cells("D4:E4")
        for row in ws.iter_rows(min_row=4, max_row=5):
            for cell in row:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if department == "purchase":
            items = [c for c in world.contracts.values() if c.side == "purchase"]
        else:
            items = [c for c in world.contracts.values() if c.department == department and c.side != "purchase" and c.in_registry
                     and (department != SCHOOL or world.parties[c.party].kind == "org")]
        items.sort(key=lambda c: (c.registry_status in ("расторгнут",), c.signed))
        total = Decimal(0)
        separated = False
        for index, contract in enumerate(items, start=1):
            if contract.registry_status == "расторгнут" and not separated:
                ws.append([])
                separated = True
            if contract.party in world.entities:
                name, number = world.entities[contract.party].name, world.entities[contract.party].number
            else:
                party = world.parties[contract.party]
                name, number = party.registry or party.name, party.registry_number
            amount = contract.registry_amount if contract.registry_amount is not None else contract.amount_on(TODAY)
            amount = text_amounts.get(contract.key, amount)
            if isinstance(amount, (int, float, Decimal)):
                total += Decimal(str(amount))
                amount = float(amount)
            executor = world.entities[contract.registry_executor or contract.entity].short
            people = ", ".join(person for person, _share in contract.staff)
            place = contract.premises if department == RENT else contract.department
            ws.append([index, name, number, contract.registry_number or contract.number, contract.signed, amount, executor, place,
                       people, contract.registry_status, contract.registry_note])
            ws.cell(ws.max_row, 5).number_format = "dd.mm.yyyy"
            if department == CLEANING and index == 3:
                for phantom in world.phantoms:
                    ws.append(["", phantom["registry"], "", phantom["number"], None, float(phantom["amount"]), "Тумар Сервис", CLEANING,
                               "", phantom["status"], "не подписан"])
                    ws.row_dimensions[ws.max_row].hidden = True
        ws.append([None, "Итого:", None, None, None, float(total)])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for column, width in zip("ABCDEFGHIJK", (5, 34, 15, 15, 12, 18, 16, 20, 22, 16, 26)):
            ws.column_dimensions[column].width = width
    return _xlsx(wb)


def rent_book(world: World) -> bytes:
    """Таблица начислений аренды: лист на месяц; «Сентябрь» забыли, в «Июле» колонки переставлены, в «Октябре» две «Сумма»."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    contracts = [c for c in world.contracts.values() if c.department == RENT and c.side != "purchase"]
    contracts.sort(key=lambda c: c.premises)
    base = ["Арендатор", "Договор", "Помещение", "Наше ТОО", "Аренда", "Коммунальные", "Итого к оплате", "Оплачено", "Примечание"]
    for month in month_range(SINCE, TODAY):
        if month.month == 9:
            continue
        ws = wb.create_sheet(MONTH_SHEETS[month.month])
        head = list(base)
        if month.month == 6:
            head.insert(4, "Долг на 01.06")
        if month.month == 7:
            head = ["Помещение", "Арендатор", "Аренда", "Коммунальные", "Договор", "Наше ТОО", "Оплачено", "Итого к оплате", "Примечание"]
        if month.month == 10:
            head = [("Сумма" if title in ("Аренда", "Коммунальные") else title) for title in head]
        ws.append([f"Начисления аренды за {MONTH_SHEETS[month.month].lower()} {month.year}"])
        ws.append([])
        ws.append(head)
        for contract in contracts:
            accrual = next((a for a in world.accruals(contract.key) if a.month == month), None)
            if accrual is None:
                continue
            if contract.party in world.entities:
                tenant = world.entities[contract.party].short
            else:
                tenant = world.parties[contract.party].registry or world.parties[contract.party].name
            paid = sum((a for p in world.payments for c, mm, a in p.alloc if c == contract.key and mm == month), Decimal(0))
            note = "съехал 30.06" if contract.key == "kainar" else ""
            values = {
                "Арендатор": tenant, "Договор": contract.number, "Помещение": contract.premises, "Наше ТОО": "Тумар Офис",
                "Аренда": float(accrual.fixed), "Коммунальные": float(accrual.utilities) or None, "Итого к оплате": float(accrual.amount),
                "Оплачено": float(paid) or None, "Примечание": note,
            }
            if month.month == 6:
                opening = world.opening(contract.key)
                values["Долг на 01.06"] = float(contract.amount_on(SINCE)) if contract.key == "kaspiy" else (float(opening) if opening else None)
            if month.month == 10:
                row = []
                for title in base[:4]:
                    row.append(values[title])
                row += [values["Аренда"], values["Коммунальные"], values["Итого к оплате"], values["Оплачено"], values["Примечание"]]
            else:
                row = [values[title] for title in head]
            ws.append(row)
    return _xlsx(wb)


def school_book(world: World) -> bytes:
    """Таблица учебного центра: слушатели с оплатами; даты текстом «дд/мм/гггг», все дни не больше 12."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Слушатели"
    ws.append(["ФИО / Организация", "ИИН/БИН", "Курс", "Начало", "Стоимость", "Оплата 1", "Дата 1", "Оплата 2", "Дата 2",
               "Оплата 3", "Дата 3", "Остаток"])
    contracts = sorted((c for c in world.contracts.values() if c.entity == "ИП"), key=lambda c: c.start)
    for contract in contracts:
        party = world.parties[contract.party]
        paid = [(p.day, a) for p in world.payments for c, _m, a in p.alloc if c == contract.key]
        paid.sort()
        price = contract.registry_amount if contract.registry_amount is not None else float(contract.amount_on(contract.start))
        row = [party.registry or party.name, party.number if party.kind == "org" else "", contract.subject, f"{contract.start:%d/%m/%Y}", price]
        for day, amount in paid[:3]:
            row += [float(amount), f"{day:%d/%m/%Y}"]
        row += [None, None] * (3 - len(paid[:3]))
        row.append(float(world.debt(contract.key)))
        ws.append(row)
    debts = wb.create_sheet("Долги на 01.06")
    debts.append(["ФИО / Организация", "Курс", "Долг на 01.06"])
    for contract in contracts:
        opening = world.opening(contract.key)
        if opening:
            debts.append([world.parties[contract.party].registry or world.parties[contract.party].name, contract.subject, float(opening)])
    return _xlsx(wb)


# ── Банк ─────────────────────────────────────────────────────────────────────


def _bic(iban: str) -> str:
    return BIC.get(iban[4:7], "")


def kaspi_statement(world: World, entity_key: str, month: date, opening: Decimal) -> tuple[bytes, Decimal]:
    """Выписка Kaspi Business по счёту юрлица за месяц - как её отдаёт банк; вторым - исходящий остаток."""
    entity = world.entities[entity_key]
    end = min(month_end(month), TODAY)
    ops = [p for p in world.payments if month <= p.day <= end and entity_key in (p.src, p.dst)]
    debit = sum((p.amount for p in ops if p.src == entity_key), Decimal(0))
    credit = sum((p.amount for p in ops if p.dst == entity_key), Decimal(0))
    closing = opening + credit - debit
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([])
    ws.append([])
    ws.append(["Текущий счет: ", None, entity.iban])
    ws.append(["Валюта счета:", None, "KZT"])
    ws.append(["Период:", None, f"{month:%d.%m.%Y} - {end:%d.%m.%Y}"])
    ws.append(["Дата последнего движения:", None, f"{(ops[-1].day if ops else end):%d.%m.%Y} 18:00"])
    ws.append(["ИИН/БИН: ", None, entity.number])
    ws.append(["Наименование: ", None, entity.name])
    ws.append(["Входящий остаток:", None, float(opening), None, "KZT"])
    ws.append(["Исходящий остаток:", None, float(closing), None, "KZT"])
    ws.append([])
    ws.append(["№\nдокумента", "Дата операции", "Дебет", "Кредит", "Наименование бенефициара / отправителя денег",
               "ИИК бенефициара / отправителя денег", "БИК банка бенефициара (отправителя денег)", "КНП", "Назначение платежа"])
    ws.append([1, 2, 3, 4, 5, 6, 7, 8, 9])
    for index, payment in enumerate(ops):
        other = payment.dst if payment.src == entity_key else payment.src
        if other in world.entities:
            them = world.entities[other]
            name, number, iban = them.name.upper(), them.number, them.iban
        else:
            party = world.parties[other]
            name = party.bank_name_on(payment.day) or party.name.upper()
            number = party.number if party.bank_number is None else party.bank_number
            iban = party.iban
        knp = {"own": "390", "deposit": "859", "refund": "859"}.get(payment.kind, "859")
        stamp = f"{payment.day:%d.%m.%Y} {10 + index % 8:02d}:{(index * 7) % 60:02d}:{(index * 13) % 60:02d}"
        out = payment.src == entity_key
        ws.append([str(1000 + index), stamp, float(payment.amount) if out else None, None if out else float(payment.amount),
                   f"{name}\nИИН/БИН {number}", iban, _bic(iban), knp, payment.purpose])
    ws.append([None, "Итого обороты в валюте счета", float(debit), float(credit)])
    ws.append([None, "Итого операций за период", str(sum(1 for p in ops if p.src == entity_key)), str(sum(1 for p in ops if p.dst == entity_key))])
    ws.append([])
    ws.append([None, f"Отчет сформирован пользователем Бухгалтер {end + timedelta(days=1):%d.%m.%Y} 10:15"])
    ws.append([None, 'Наименование и БИК обслуживающего Банка: АО "KASPI BANK" Бик: CASPKZKA'])
    return _xlsx(wb), closing


#: С какого месяца у юрлица выписки: у «Тумар Сервис» - с мая (предоплата за июнь, О9).
STATEMENTS_FROM = {"ТС": date(2026, 5, 1), "ТО": SINCE, "ТП": date(2026, 8, 1), "ИП": SINCE}
OPENING_CASH = {"ТС": Decimal(6_500_000), "ТО": Decimal(3_200_000), "ТП": Decimal(0), "ИП": Decimal(450_000)}


def statements(world: World) -> dict[str, bytes]:
    out = {}
    for key, since in STATEMENTS_FROM.items():
        balance = OPENING_CASH[key]
        for month in month_range(since, TODAY):
            data, balance = kaspi_statement(world, key, month, balance)
            out[f"Банк/Выписка Kaspi {world.entities[key].short} {month:%Y-%m}.xlsx"] = data
    return out


# ── Правда и описание ────────────────────────────────────────────────────────


def _money(value: Decimal) -> str:
    return format(value, "f")


def truth(world: World) -> dict[str, Any]:
    """Правда для проверки: участники, договоры, начисления, оплаты, сальдо на 01.06 и долги на 25.10."""
    return {
        "today": TODAY.isoformat(),
        "since": SINCE.isoformat(),
        "entities": {k: {"name": e.name, "number": e.number, "iban": e.iban, "kind": e.kind, "vat_from": e.vat_from and e.vat_from.isoformat(),
                         "regime": e.regime, "onec": e.onec, "created": e.created.isoformat()} for k, e in world.entities.items()},
        "parties": {k: {"number": p.number, "kind": p.kind, "name": p.name, "onec": p.onec_name, "registry": p.registry, "bank": p.bank,
                        "registry_number": p.registry_number, "renamed": p.renamed and [p.renamed[0].isoformat(), p.renamed[1]],
                        "cases": list(p.cases)} for k, p in world.parties.items()},
        "contracts": {
            k: {
                "entity": c.entity, "party": c.party, "number": c.number, "signed": c.signed.isoformat(), "start": c.start.isoformat(),
                "end": c.end and c.end.isoformat(), "billing": c.billing, "side": c.side, "department": c.department,
                "amounts": [[d.isoformat(), _money(a)] for d, a in c.amounts],
                "executors": [[d.isoformat(), code] for d, code in c.executors],
                "pauses": [[a.isoformat(), b.isoformat()] for a, b in c.pauses],
                "deposit": _money(c.deposit), "in_registry": c.in_registry, "cases": list(c.cases),
                "accruals": [[a.month.isoformat(), _money(a.amount), a.executor] for a in world.accruals(k)],
                "opening": _money(world.opening(k)), "debt": _money(world.debt(k)),
            }
            for k, c in world.contracts.items()
        },
        "payments": [
            {"key": p.key, "day": p.day.isoformat(), "src": p.src, "dst": p.dst, "amount": _money(p.amount), "kind": p.kind,
             "purpose": p.purpose, "alloc": [[c, m and m.isoformat(), _money(a)] for c, m, a in p.alloc], "cases": list(p.cases)}
            for p in world.payments
        ],
        "owed": {k: _money(v) for k, v in sorted(world.owed().items())},
        "revenue_2026": {k: _money(v) for k, v in sorted(world.revenue().items())},
        "cases": {case.id: {"group": case.group, "title": case.title, "expect": case.expect, "answer": case.answer, "needs": case.needs,
                            "where": case_index(world).get(case.id, [])} for case in CASES},
    }


def readme(world: World) -> str:
    owed = world.owed()
    lines = [
        "# Полигон «Тумар»",
        "",
        "Выдуманная группа компаний для проверки KORT без Google-таблиц. Сгенерировано `tests/polygon` (seed по умолчанию).",
        "",
        f"Сегодня полигона - {TODAY:%d.%m.%Y}, учёт оплат в KORT - с {SINCE:%d.%m.%Y}. Правда - `truth.json`.",
        "",
        "## Юрлица",
        "",
    ]
    for entity in world.entities.values():
        source = {"online": "1С в интернете (снимок OData)", "files": "1С на компьютере (файлы)", "none": "без 1С, Excel"}[entity.onec]
        lines.append(f"- {entity.name} ({entity.number}) - {source}; нам должны на {TODAY:%d.%m}: {owed.get(entity.key, 0):,.0f} ₸".replace(",", " "))
    lines += ["", "## Каталог неоднозначностей", "", "| Код | Что в данных | KORT | Ответ человека | Не хватает ядру |", "| --- | --- | --- | --- | --- |"]
    for case in CASES:
        lines.append(f"| {case.id} | {case.title} | {case.expect} | {case.answer} | {case.needs or '-'} |")
    return "\n".join(lines) + "\n"


def render(world: World) -> dict[str, Any]:
    files: dict[str, Any] = {
        "1С Тумар Сервис/odata.json": onec_odata(world, "ТС"),
        "1С Тумар Сервис/ОСВ 1210 июнь 2026.xlsx": osv_1210(world, ("ТС",)),
        "1С Тумар Офис/Контрагенты.xlsx": onec_parties(world, ("ТО", "ТП")),
        "1С Тумар Офис/Реализация товаров и услуг 2026.xlsx": onec_journal(world, ("ТО", "ТП"), "act"),
        "1С Тумар Офис/Счета на оплату 2026.xlsx": onec_journal(world, ("ТО", "ТП"), "invoice"),
        "1С Тумар Офис/ОСВ 1210 июнь 2026.xlsx": osv_1210(world, ("ТО", "ТП")),
        "Excel/Реестр договоров 2026.xlsx": registry(world),
        "Excel/Аренда - начисления 2026.xlsx": rent_book(world),
        "Excel/Учебный центр 2026.xlsx": school_book(world),
        **statements(world),
        "truth.json": truth(world),
        "README.md": readme(world),
    }
    return files


def write(world: World, folder: str) -> list[str]:
    import pathlib

    root = pathlib.Path(folder)
    written = []
    for name, content in render(world).items():
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        if isinstance(content, (dict, list)):
            path.write_text(json.dumps(content, ensure_ascii=False, indent=1), encoding="utf-8")
        elif isinstance(content, str):
            path.write_text(content, encoding="utf-8")
        else:
            path.write_bytes(content)
        written.append(name)
    return written


__all__ = ["kaspi_statement", "onec_odata", "osv_1210", "ref", "registry", "render", "rent_book", "school_book", "truth", "write"]
