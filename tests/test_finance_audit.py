"""Журнал действий пишется всегда: обход всех изменяющих маршрутов «Финансов».

Один сценарий проходит каждый POST / PATCH / PUT / DELETE раздела - учёт,
договоры, люди, права, входы - и после каждого проверяет, что в журнале
появилось событие с автором, сеансом, адресом, браузером и видом. Маршрут,
который ничего не пишет, стоит в `ALLOWED` с причиной; новый маршрут без
события и без причины роняет тест.

Почему обход, а не проверка по месту: правка ячейки в листе журнала меняла
операцию молча - событие писали карточка и импорт, а лист забыли, и
«кто поменял сумму» по нему узнать было нельзя. Так забывают всегда в
одном месте; ловит это только полный перебор.
"""
from __future__ import annotations

import io
import json
import uuid
from typing import Any

import pytest
import sqlalchemy as sa
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.finance import audit, auth
from app.finance.db import finance_session
from app.finance.models import ACTION_CATEGORIES, ActionLog

BASE = "/api/v1/finance"
PASSWORD = "pass-12345"
AGENT = "pytest-agent"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _sales_file() -> bytes:
    """Журнал «Реализация товаров и услуг» 1С, выгруженный списком: одна реализация."""
    from datetime import datetime

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Реализации"
    sheet.append(["Дата", "Номер", "Организация", "Контрагент", "Договор", "Сумма", "Валюта", "Содержание"])
    sheet.append([datetime(2026, 9, 30, 18), "00000000001", "ТОО BBC", "ТОО Альфа", "Договор № 1 от 01.01.2026", 100000, "KZT",
                  "Услуги за период с 01.09.2026 по 30.09.2026"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _listeners_file() -> bytes:
    """Таблица отдела, которую KORT сам не узнаёт: слушатели курсов."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Слушатели"
    sheet.append(["ФИО / Организация", "Курс", "Начало", "Стоимость"])
    sheet.append(["Касымов Ерлан", "Бухгалтерия с нуля", "15.09.2026", 350000])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _osv_file() -> bytes:
    """ОСВ по 1210 в форме отчёта 1С: шапка в две строки, контрагент и его договор отступом."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    book = Workbook()
    sheet = book.active
    sheet.append(["ТОО BBC"])
    sheet.append(["Оборотно-сальдовая ведомость по счету 1210 за Сентябрь 2026 г."])
    sheet.append([])
    sheet.append(["Счет", "Сальдо на начало периода", None, "Обороты за период", None, "Сальдо на конец периода", None])
    sheet.append([None, "Дебет", "Кредит", "Дебет", "Кредит", "Дебет", "Кредит"])
    sheet.merge_cells("B4:C4")
    sheet.merge_cells("D4:E4")
    sheet.merge_cells("F4:G4")
    sheet.append(["1210", 100000, None, None, None, 100000, None])
    sheet.append(["ТОО Альфа", 100000, None, None, None, 100000, None])
    sheet.cell(sheet.max_row, 1).alignment = Alignment(indent=1)
    sheet.append(["Договор № 1 от 01.01.2026", 100000, None, None, None, 100000, None])
    sheet.cell(sheet.max_row, 1).alignment = Alignment(indent=2)
    sheet.append(["Итого", 100000, None, None, None, 100000, None])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()

#: Изменяющие маршруты без события - и почему.
ALLOWED: dict[tuple[str, str], str] = {
    ("POST", "/auth/phone/start"): "шаг входа: только отвечает, какой экран показать, ничего не меняет",
    ("POST", "/contracts/{contract_id}/amendments/parse"): "читает текст соглашений, в базу не пишет",
    ("POST", "/contracts/setup/views/preview"): "считает договоры под правило листа, в базу не пишет",
    ("POST", "/core/vat/preview"): "считает оборот юрлиц до и после перевода клиентов, в базу не пишет",
    ("POST", "/core/intake"): "приёмник называет форму каждого листа файла, в базу не пишет",
    ("POST", "/core/intake/plan"): "предложение и сверка таблицы по разметке, в базу не пишет",
    ("POST", "/contracts/imports/{batch_id}/decide"): (
        "решения протокола - черновик разбора; события пишут загрузка файла и заведение"
    ),
    ("PATCH", "/import/batches/{batch_id}/rows/{line}"): (
        "правка черновой строки партии: в учёт не попадает, событие пишет заведение партии"
    ),
    ("POST", "/sheets/preview"): (
        "нужна книга Google; ответ собирает тот же `_preview_response`, что у /import/preview, "
        "и его событие проверено там"
    ),
    ("POST", "/recurrences/materialize"): (
        "пустое продление ничего не меняет и не пишет; с новыми ожиданиями пишет recurrence.materialize"
    ),
    ("POST", "/bills/template"): (
        "разбирает загруженный шаблон счетов и возвращает его строки; в книги каждую строку пишет "
        "POST /bills/{bill_id}, и событие - там"
    ),
    ("PUT", "/looks/{key}"): "личный вид листа - настройка экрана одной учётки, не данные компании",
    ("DELETE", "/looks/{key}"): "сброс личного вида листа - настройка экрана одной учётки, не данные компании",
    ("POST", "/looks/habits/{group}"): "привычка вида раздела (минуты в «Таблице» и «Карточках») - экран одной учётки",
    ("POST", "/support/messages"): (
        "переписка учётки с поддержкой KORT - разговор с разработчиком, а не действие в учёте; след - в канале поддержки"
    ),
    ("POST", "/support/read"): "отметка «прочитано» у ответов поддержки - экран одной учётки",
    ("POST", "/support/notice"): "объявление о поддержке показано или выключено - экран одной учётки",
    ("POST", "/auth/device/claim"): (
        "забор сеанса на новом устройстве; событие входа пишет `_begin`, то же, что у /auth/login; "
        "сам забор проверяется в test_finance_devices.py"
    ),
    ("POST", "/auth/device/{device_id}/approve"): (
        "подтверждение входа с нового устройства - своё решение учётки о своём же входе, не действие в учёте компании"
    ),
    ("POST", "/auth/device/{device_id}/deny"): (
        "отклонение входа с нового устройства - своё решение учётки, не действие в учёте компании"
    ),
}


@pytest.fixture
def finance_db(tmp_path, monkeypatch):
    from sqlalchemy import create_engine

    from app.core.config import settings
    from app.finance import db as finance_db_module

    engine = create_engine(f"sqlite:///{tmp_path / 'finance.db'}", future=True)
    monkeypatch.setattr(finance_db_module, "get_engine", lambda: engine)
    monkeypatch.setattr(finance_db_module, "_session_factory", None)
    monkeypatch.setattr(finance_db_module, "_initialized", False)
    monkeypatch.setattr(settings, "environment", "test")
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()
    yield engine
    auth._ATTEMPTS.clear()
    audit._VIEWS.clear()


@pytest.fixture
def app(finance_db) -> FastAPI:
    from finance_routes import finance_app

    return finance_app()


def client(app: FastAPI, ip: str = "10.0.0.7") -> TestClient:
    return TestClient(app, headers={"x-client-ip": ip, "user-agent": AGENT})


class Walk:
    """Вызов маршрута и проверка его события."""

    def __init__(self) -> None:
        self.covered: set[tuple[str, str]] = set()
        self.owner_id: uuid.UUID | None = None

    @staticmethod
    def _ids() -> set[uuid.UUID]:
        with finance_session() as session:
            return set(session.scalars(sa.select(ActionLog.id)))

    @staticmethod
    def _rows(ids: set[uuid.UUID]) -> list[dict[str, Any]]:
        if not ids:
            return []
        with finance_session() as session:
            return [
                {
                    "kind": row.kind, "category": row.category, "user_id": row.user_id,
                    "session_id": row.session_id, "ip": row.ip, "user_agent": row.user_agent,
                }
                for row in session.scalars(sa.select(ActionLog).where(ActionLog.id.in_(ids)))
            ]

    def call(
        self,
        who: TestClient,
        method: str,
        template: str,
        *,
        user: Any = "owner",
        session: bool = True,
        ip: str = "10.0.0.7",
        agent: str = AGENT,
        **kwargs: Any,
    ):
        params = {key: kwargs.pop(key) for key in list(kwargs) if "{" + key + "}" in template}
        path = BASE + template.format(**params)
        before = self._ids()
        response = who.request(method, path, **kwargs)
        key = (method, template)
        assert response.status_code < 400, f"{method} {template}: {response.status_code} {response.text}"
        rows = self._rows(self._ids() - before)
        assert rows, f"{method} {template}: событие в журнал не записано"
        expected = self.owner_id if user == "owner" else user
        good = [
            row for row in rows
            if (row["user_id"] is not None if expected == "any" else row["user_id"] == expected)
            and (row["session_id"] is not None) == session
            and row["ip"] == ip
            and row["user_agent"] == agent
            and row["category"] in ACTION_CATEGORIES
        ]
        assert good, f"{method} {template}: у события не тот автор, сеанс или адрес: {rows}"
        self.covered.add(key)
        return response


def _journal_file() -> bytes:
    import openpyxl
    from datetime import date

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Дата платежа", "Сумма", "Со счёта", "На счёт", "Категория", "Контрагент", "Комментарий"])
    sheet.append([date(2026, 8, 3), 150000, None, "Касса", "Выручка", "Клиент", "оплата"])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _section(batch: dict, key: str) -> dict:
    return next(section for section in batch["report"]["sections"] if section["key"] == key)


def test_kazhdyy_izmenyayushchiy_marshrut_pishet_sobytie(app: FastAPI, monkeypatch: pytest.MonkeyPatch) -> None:
    from finance_routes import api_routes
    from test_finance_contracts_import import _registry_file

    walk = Walk()
    owner = client(app)

    # ── регистрация ─────────────────────────────────────────────────────
    me = walk.call(
        owner, "POST", "/auth/register", user="any",
        json={"email": "owner@bbc.kz", "password": PASSWORD, "company": "BBC (тест)", "full_name": "Ермеков Нурболат"},
    ).json()
    walk.owner_id = uuid.UUID(me["user"]["id"])
    first_company = me["company"]["id"]

    # ── счета и справочники ────────────────────────────────────────────
    accounts = {item["name"]: item["id"] for item in owner.get(f"{BASE}/dictionaries").json()["accounts"]}
    cash = accounts["Касса"]
    reserve = walk.call(owner, "POST", "/accounts", json={"name": "Резерв"}).json()["id"]
    walk.call(owner, "PATCH", "/accounts/{account_id}", account_id=reserve, json={"starting_balance": "1000"})
    walk.call(owner, "PUT", "/accounts/{account_id}/number", account_id=reserve, json={"number": "KZ86125KZT5004100100"})
    category = walk.call(
        owner, "POST", "/dictionaries/{kind}", kind="categories", json={"name": "Аренда офиса", "side": "expense"}
    ).json()["id"]
    walk.call(owner, "PATCH", "/dictionaries/categories/{item_id}/nature", item_id=category, json={"nature": "operating"})
    tag = owner.post(f"{BASE}/dictionaries/tags", json={"name": "срочно"}).json()["id"]
    walk.call(owner, "DELETE", "/dictionaries/{kind}/{item_id}", kind="tags", item_id=tag)

    # ── журнал и лист ───────────────────────────────────────────────────
    operation = walk.call(
        owner, "POST", "/operations",
        json={"kind": "income", "paid_at": "2026-09-01", "amount": "1000", "account_to_id": cash, "comment": "аренда"},
    ).json()["id"]
    walk.call(owner, "PATCH", "/operations/{operation_id}", operation_id=operation, json={"comment": "аренда сентябрь"})
    walk.call(
        owner, "PATCH", "/grid/cell",
        json={"operation_id": operation, "column": "comment", "value": "аренда · из листа"},
    )
    grid_row = walk.call(
        owner, "POST", "/grid/row",
        json={"cells": {"paid_at": "02.09.2026", "amount": "500", "account_to": "Касса", "comment": "строка листа"}},
    ).json()["id"]
    # Правка из листа откатывается, как правка карточки.
    cell_edit = next(
        item for item in owner.get(f"{BASE}/audit", params={"kind": "operation.update"}).json()["items"]
        if item["title"].startswith("правка в таблице")
    )
    assert cell_edit["can_undo"] and cell_edit["before"]["comment"] == "аренда сентябрь"
    walk.call(owner, "POST", "/history/{entry_id}/undo", entry_id=cell_edit["id"])
    created = next(
        item for item in owner.get(f"{BASE}/audit", params={"kind": "operation.create", "entity_id": grid_row}).json()["items"]
    )
    walk.call(owner, "POST", "/audit/{entry_id}/undo", entry_id=created["id"])
    walk.call(owner, "DELETE", "/operations/{operation_id}", operation_id=operation)

    # ── планы, правила, разметка ────────────────────────────────────────
    walk.call(owner, "POST", "/plans", json={"month": "2026-09-01", "side": "income", "amount": "100000"})
    rule = walk.call(
        owner, "POST", "/rules",
        json={"name": "Аренда", "conditions": [{"field": "comment", "op": "contains", "value": "аренда"}],
              "actions": {"category": "Аренда офиса"}},
    ).json()["id"]
    walk.call(owner, "PATCH", "/rules/{rule_id}", rule_id=rule, json={"active": False})
    walk.call(owner, "POST", "/rules/apply", json={})
    walk.call(owner, "DELETE", "/rules/{rule_id}", rule_id=rule)
    walk.call(owner, "POST", "/autotag", json={"groups": [{"side": "expense", "category": "Аренда"}]})

    # ── загрузка выписки ────────────────────────────────────────────────
    batch = walk.call(
        owner, "POST", "/import/preview", files={"file": ("выписка.xlsx", _journal_file(), XLSX)}
    ).json()["batch_id"]
    walk.call(owner, "POST", "/import/batches/{batch_id}/apply", batch_id=batch, json={})

    # ── счета на оплату, повторения ─────────────────────────────────────
    invoice = walk.call(
        owner, "POST", "/invoices",
        json={"issued_at": "2026-09-01", "due_at": "2026-09-15", "lines": [{"title": "Услуга", "price": "1000"}]},
    ).json()["id"]
    walk.call(owner, "POST", "/invoices/{invoice_id}/void", invoice_id=invoice)
    recurrence = walk.call(
        owner, "POST", "/recurrences",
        json={"title": "Аренда", "kind": "expense", "amount": "1000", "start_at": "2026-09-01", "account_from_id": cash},
    ).json()["id"]
    walk.call(owner, "PATCH", "/recurrences/{recurrence_id}", recurrence_id=recurrence, params={"active": "false"})
    walk.call(owner, "DELETE", "/recurrences/{recurrence_id}", recurrence_id=recurrence)

    # ── подключения ─────────────────────────────────────────────────────
    integration = walk.call(
        owner, "POST", "/integrations", json={"slug": "kaspi", "kind": "api", "title": "Kaspi", "account_id": cash}
    ).json()
    token = walk.call(owner, "POST", "/integrations/{integration_id}/token", integration_id=integration["id"]).json()["token"]
    walk.call(owner, "PATCH", "/integrations/{integration_id}", integration_id=integration["id"], params={"state": "active"})
    # Приём по токену: автора-человека нет, подпись - подключение.
    walk.call(
        client(app), "POST", "/integrations/inbox", user=None, session=False, ip="", agent="",
        headers={"x-finance-token": token},
        json={"operations": [{"paid_at": "2026-09-03", "amount": "700", "comment": "из банка"}]},
    )
    walk.call(owner, "DELETE", "/integrations/{integration_id}", integration_id=integration["id"])

    # ── базы 1С: REST-сервис подменён - сеть в тесте не нужна ───────────
    from app.finance.config import finance_settings
    from app.finance.onec import service as onec_service
    from test_finance_onec import _fake_1c

    monkeypatch.setattr(finance_settings, "secrets_key", "audit-key")
    monkeypatch.setattr(onec_service, "check_host", lambda url: None)
    fake_1c = _fake_1c()
    monkeypatch.setattr(onec_service, "Client", lambda connection: fake_1c)
    onec_base = walk.call(
        owner, "POST", "/onec/bases",
        json={"url": "https://1cfresh.kz/a/ea201/284171", "login": "odata.user", "password": "пароль"},
    ).json()["id"]
    walk.call(owner, "PATCH", "/onec/bases/{base_id}", base_id=onec_base, json={"title": "ИБ ТОО BBC"})
    walk.call(owner, "POST", "/onec/bases/{base_id}/sync", base_id=onec_base)
    loose = owner.get(f"{BASE}/onec/month", params={"month": "2026-09"}).json()["unmatched"][0]["id"]
    walk.call(owner, "POST", "/onec/documents/{document_id}/decide", document_id=loose, json={"action": "none"})
    # Вход из 1С: ОСВ к базе, наше юрлицо базы и оплаты юрлица из 1С.
    walk.call(
        owner, "POST", "/core/onec-start/{base_id}/osv", base_id=onec_base,
        files={"file": ("ОСВ 1210.xlsx", _osv_file(), XLSX)},
    )
    assert owner.post(f"{BASE}/core/onec-start/{onec_base}/entity", json={"answers": {}}).status_code == 200
    walk.call(owner, "POST", "/core/onec-start/{base_id}/{step}", base_id=onec_base, step="payments", json={"answers": {}})
    # Сальдо 1210 из 1С по REST: чтение 1С подменено готовой ведомостью.
    from datetime import date as _date
    from decimal import Decimal as _Decimal

    from app.finance.onec import balances as onec_balances
    from app.finance.onec.osv import Line, Osv

    rest = Osv("1С по REST", "", _date(2026, 9, 1), _date(2026, 9, 30), "1210", lines=[
        Line("ТОО Альфа", "Договор № 1 от 01.01.2026", "1", _date(2026, 1, 1), _Decimal(100000), _Decimal(0), _Decimal(0), _Decimal(0), 1),
    ], opening_total=_Decimal(100000))
    monkeypatch.setattr(onec_balances, "read", lambda *args, **kwargs: rest)
    walk.call(owner, "POST", "/core/onec-start/{base_id}/balances", base_id=onec_base, json={"day": "2026-09-01"})
    walk.call(owner, "DELETE", "/onec/bases/{base_id}", base_id=onec_base)
    # 1С без интернета: журнал реализаций файлом - база организации из файлов.
    walk.call(
        owner, "POST", "/core/onec-files",
        files=[("files", ("Реализация товаров и услуг.xlsx", _sales_file(), XLSX))],
    )
    # Excel любой формы: разметка таблицы отдела запоминается словарём компании.
    listeners = json.dumps({"sheets": {"Слушатели": {
        "form": "contract", "columns": {"ФИО / Организация": "customer", "Курс": "subject", "Начало": "billing_from", "Стоимость": "amount"},
        "billing": "total",
    }}})
    owner.post(f"{BASE}/core/intake/plan", files={"file": ("Учебный центр.xlsx", _listeners_file(), XLSX)}, data={"answers": listeners})
    walk.call(owner, "POST", "/core/intake/apply", files={"file": ("Учебный центр.xlsx", _listeners_file(), XLSX)},
              data={"answers": listeners})

    # ── компании ────────────────────────────────────────────────────────
    second = walk.call(owner, "POST", "/auth/companies", json={"title": "Вторая"}).json()["id"]
    walk.call(owner, "POST", "/auth/switch", json={"company_id": second})
    owner.post(f"{BASE}/auth/switch", json={"company_id": first_company})
    walk.call(owner, "PATCH", "/auth/company", json={"title": "BBC (тест)"})

    # ── приглашение по почте (прежние роли) ─────────────────────────────
    invited = walk.call(
        owner, "POST", "/auth/members",
        json={"email": "buh@bbc.kz", "password": "temp-pass-1", "role": "accountant", "full_name": "Бухгалтер Б"},
    ).json()["id"]
    # «Забыли?» у входа по почте - просьба администраторам, автор неизвестен.
    walk.call(client(app), "POST", "/auth/forgot", user=None, session=False, json={"email": "buh@bbc.kz"})
    walk.call(owner, "PATCH", "/auth/members/{user_id}", user_id=invited, json={"role": "viewer"})
    walk.call(owner, "DELETE", "/auth/members/{user_id}", user_id=invited)

    # ── свои сеансы, пароль, профиль ────────────────────────────────────
    other = client(app)
    walk.call(other, "POST", "/auth/login", json={"email": "owner@bbc.kz", "password": PASSWORD})
    foreign = next(item for item in owner.get(f"{BASE}/auth/sessions").json()["items"] if not item["current"])
    walk.call(owner, "DELETE", "/auth/sessions/{session_id}", session_id=foreign["id"])
    client(app).post(f"{BASE}/auth/login", json={"email": "owner@bbc.kz", "password": PASSWORD})
    walk.call(owner, "POST", "/auth/sessions/end-others")
    walk.call(owner, "POST", "/auth/password", json={"old_password": PASSWORD, "new_password": "new-pass-123"})
    walk.call(owner, "POST", "/auth/email", json={"email": "boss@bbc.kz", "password": "new-pass-123"})
    assert owner.post(f"{BASE}/auth/email", json={"email": "owner@bbc.kz", "password": "new-pass-123"}).status_code == 200
    walk.call(owner, "PATCH", "/auth/profile", json={"phone": "+77770001122"})
    leaving = client(app)
    leaving.post(f"{BASE}/auth/login", json={"email": "owner@bbc.kz", "password": "new-pass-123"})
    walk.call(leaving, "POST", "/auth/logout")

    # ── люди, права, вход по номеру ─────────────────────────────────────
    department = walk.call(owner, "POST", "/people/departments", json={"code": "ЮО", "title": "Юридический отдел"}).json()["id"]
    walk.call(owner, "PATCH", "/people/departments/{department_id}", department_id=department, json={"title": "Юристы"})
    person = walk.call(
        owner, "POST", "/people/employees",
        json={"full_name": "Сейтова Айдана", "phone": "+77025550122", "department_id": department, "access": True},
    ).json()
    person_user = uuid.UUID(person["account"]["user_id"])
    walk.call(owner, "PATCH", "/people/employees/{employee_id}", employee_id=person["id"], json={"job_title": "Юрист"})
    walk.call(owner, "PUT", "/access/{kind}/{subject_id}", kind="employee", subject_id=person["id"],
              json={"changes": {"journal": "view"}})
    phone = client(app, "10.0.0.9")
    assert phone.post(f"{BASE}/auth/phone/start", json={"phone": "+77025550122"}).json() == {"step": "set_password"}
    walk.call(phone, "POST", "/auth/phone/set-password", user=person_user, session=False, ip="10.0.0.9",
              json={"phone": "+77025550122", "password": "secret-123"})
    walk.call(phone, "POST", "/auth/phone/login", user=person_user, ip="10.0.0.9",
              json={"phone": "+77025550122", "password": "secret-123"})
    walk.call(phone, "POST", "/auth/phone/forgot", user=None, session=False, ip="10.0.0.9",
              json={"phone": "+77025550122"})
    request = next(item for item in owner.get(f"{BASE}/notifications").json()["items"] if item["actionable"])
    walk.call(owner, "POST", "/notifications/{notification_id}/resolve", notification_id=request["id"])
    phone_session = owner.get(f"{BASE}/people/employees/{person['id']}/sessions").json()["items"][0]["id"]
    walk.call(owner, "DELETE", "/people/employees/{employee_id}/sessions/{session_id}",
              employee_id=person["id"], session_id=phone_session)
    walk.call(owner, "POST", "/people/employees/{employee_id}/end-sessions", employee_id=person["id"])
    walk.call(owner, "POST", "/people/employees/{employee_id}/block", employee_id=person["id"])
    walk.call(owner, "POST", "/people/employees/{employee_id}/unblock", employee_id=person["id"])
    walk.call(owner, "POST", "/people/employees/{employee_id}/reset", employee_id=person["id"])
    helper = owner.post(f"{BASE}/people/employees", json={"full_name": "Помощник Юриста"}).json()
    walk.call(owner, "POST", "/people/employees/{employee_id}/account", employee_id=helper["id"],
              json={"phone": "+77025550123"})
    walk.call(owner, "DELETE", "/people/employees/{employee_id}", employee_id=helper["id"])
    walk.call(owner, "POST", "/audit/view", json={"section": "journal"})

    # ── договоры ────────────────────────────────────────────────────────
    def contract_values(customer: str, number: str, **extra: Any) -> dict:
        return {"values": {"executor": "BBC", "customer": customer, "number": number, "amount": "500000", **extra}}

    first = walk.call(owner, "POST", "/contracts", json=contract_values("ТОО Альфа", "ЮО/141")).json()["contract"]["id"]
    twin = owner.post(f"{BASE}/contracts", json=contract_values("ТОО Бета", "ЮО/141")).json()["contract"]["id"]
    walk.call(owner, "PATCH", "/contracts/{contract_id}", contract_id=first, json={"values": {"note": "звонили"}})
    walk.call(owner, "POST", "/contracts/{contract_id}/acknowledge", contract_id=twin, json={"code": "number_taken"})
    customer = owner.get(f"{BASE}/contracts/{first}").json()["contract"]["values"]["customer"]
    payment = owner.post(
        f"{BASE}/operations",
        json={"kind": "income", "paid_at": "2026-09-02", "amount": "500000", "account_to_id": cash,
              "counterparty_id": customer},
    ).json()["id"]
    walk.call(owner, "POST", "/contracts/{contract_id}/payments", contract_id=first,
              json={"operation_id": payment, "action": "link"})
    dated = owner.patch(
        f"{BASE}/contracts/{first}",
        json={"values": {"amount": "750000"}, "mode": {"kind": "from_date", "effective_from": "2026-01-01"}},
    )
    assert dated.status_code == 200, dated.text
    amendment = owner.get(f"{BASE}/contracts/{first}/amendments").json()["items"][0]["id"]
    walk.call(owner, "DELETE", "/contracts/{contract_id}/amendments/{amendment_id}",
              contract_id=first, amendment_id=amendment)
    with_text = owner.post(
        f"{BASE}/contracts", json=contract_values("ТОО Гамма", "ЮО/150", amendments_text="№ 1 от 23.07.2024")
    ).json()["contract"]["id"]
    pieces = owner.post(f"{BASE}/contracts/{with_text}/amendments/parse").json()["pieces"]
    walk.call(owner, "POST", "/contracts/{contract_id}/amendments/confirm", contract_id=with_text,
              json={"piece": pieces[0]})
    walk.call(owner, "POST", "/contracts/sync", json={
        "source": "проверка", "apply": True,
        "rows": [{"number": "ЮО/141", "customer": "ТОО Альфа", "values": {"Примечания": "из книги"}}],
    })
    # Доли сотрудников и отделов, изменение листа и возврат по точке (29.09.2026).
    owner.patch(f"{BASE}/contracts/{first}", json={"values": {"people": "Сейтова Айдана"}})
    walk.call(owner, "PUT", "/contracts/{contract_id}/shares/people", contract_id=first,
              json={"unit": "percent", "items": [{"employee_id": person["id"], "value": "100"}]})
    walk.call(owner, "PUT", "/contracts/{contract_id}/shares/departments", contract_id=first,
              json={"unit": "percent", "items": [{"department_id": department, "value": "100"}]})
    moved = walk.call(owner, "POST", "/contracts/sheet/change",
                      json={"action": "move_rows", "ids": [with_text], "before": first, "view": "main"}).json()
    walk.call(owner, "POST", "/contracts/restore-points/{point_id}/restore", point_id=moved["point"]["id"])
    walk.call(owner, "DELETE", "/contracts/{contract_id}", contract_id=twin)
    # Корзина: удалённый договор вернуть, снова удалить и стереть насовсем.
    listed = owner.get(f"{BASE}/trash")
    assert listed.status_code == 200, listed.text
    assert any(item["id"] == twin for item in listed.json()["items"])
    walk.call(owner, "POST", "/trash/{kind}/{item_id}/restore", kind="contract", item_id=twin)
    owner.delete(f"{BASE}/contracts/{twin}")
    walk.call(owner, "DELETE", "/trash/{kind}/{item_id}", kind="contract", item_id=twin)

    # ── настройка реестра ───────────────────────────────────────────────
    field = walk.call(owner, "POST", "/contracts/setup/fields", json={"title": "Ссылка на Битрикс", "type": "url"}).json()["key"]
    walk.call(owner, "PATCH", "/contracts/setup/fields/{key}", key=field, json={"title": "Битрикс"})
    paused = walk.call(owner, "POST", "/contracts/setup/lists/{field_key}", field_key="status",
                       json={"value": "Приостановлен"}).json()["id"]
    pause = owner.post(f"{BASE}/contracts/setup/lists/status", json={"value": "На паузе"}).json()["id"]
    walk.call(owner, "PATCH", "/contracts/setup/values/{value_id}", value_id=paused, json={"value": "Приостановлен."})
    walk.call(owner, "POST", "/contracts/setup/values/merge", json={"keep": paused, "drop": pause})
    entity = walk.call(owner, "POST", "/contracts/setup/entities", json={"name": "BBCL", "code": "BBCL"}).json()["id"]
    walk.call(owner, "PATCH", "/contracts/setup/entities/{party_id}", party_id=entity, json={"full_name": "ТОО BBC Legal"})
    parties = {item["name"]: item["id"] for item in owner.get(f"{BASE}/contracts/parties", params={"q": "ТОО"}).json()["parties"]}
    walk.call(owner, "POST", "/contracts/setup/parties/merge", json={"keep": parties["ТОО Альфа"], "drop": parties["ТОО Бета"]})
    registry_department = walk.call(owner, "POST", "/contracts/setup/departments", json={"code": "НО"}).json()["id"]
    walk.call(owner, "PATCH", "/contracts/setup/departments/{department_id}", department_id=registry_department,
              json={"title": "Налоговый отдел"})
    rule_filter = {"any": [{"all": [{"field": "number", "op": "contains", "value": "ЮО"}]}]}
    assert owner.post(f"{BASE}/contracts/setup/views/preview", json={"filter": rule_filter}).status_code == 200
    view = walk.call(owner, "POST", "/contracts/setup/views",
                     json={"title": "Юристы", "blocks": [{"filter": rule_filter}]}).json()["id"]
    walk.call(owner, "PATCH", "/contracts/setup/views/{view_id}", view_id=view, json={"title": "Юристы ЮО"})

    # ── сводка оплат: книга Google подменена - сеть в тесте не нужна ────
    from app.finance.contracts import summary as summary_module

    monkeypatch.setattr(
        summary_module,
        "_read",
        lambda spreadsheet_id, tab, extra=None: summary_module.build_index(
            [["Заказчик (Название Фирмы)", "№ Договора", "Сумма Факт Поступ."], ["ТОО Альфа", "№ 1", "1000"]],
            title="Сводка",
            worksheet=tab,
            columns=summary_module.columns_for(extra),
        ),
    )
    walk.call(owner, "POST", "/contracts/setup/summary",
              json={"link": "https://docs.google.com/spreadsheets/d/abc123/edit", "worksheet": "Сводка"})

    # ── карточка счёта пишет в книгу юрлица: книги - в памяти ───────────
    from app.finance.contracts import bill_book
    from app.finance.contracts import bills as bills_module
    from bill_book_fake import fake_api

    monkeypatch.setattr(
        bills_module,
        "_read",
        lambda spreadsheet_id, tab, extra=None: bills_module.build_board(
            [
                ["Наша Фирма", "Заказчик (Название Фирмы)", "№ Договора", "Техн. 1", "Счет", "Период с (нач.)"],
                ["BBCL", "ТОО KENESARY AGRO", "№ЮО/44", "ОКТЯБРЬ 2026", "Сегодня день выставления", "22.10.2026"],
            ],
            title="Сводка",
            worksheet=tab,
            columns=bills_module.columns_for(extra),
        ),
    )
    books = fake_api()
    books.summary_id = "abc123"
    monkeypatch.setattr(bill_book, "_api", books)
    bill_book.forget()
    bills_module._slots.clear()  # чтение до подмены (без кредов Google) держалось бы минуту
    listed = owner.get(f"{BASE}/bills").json()
    assert len(listed["items"]) == 1, "ошибка: " + str(listed["source"]["error"])
    [bill] = listed["items"]
    card = owner.get(f"{BASE}/bills/{bill['id']}").json()
    walk.call(owner, "POST", "/bills/{bill_id}", bill_id=bill["id"],
              json={"values": {"bill_no": "312", "bill_date": "2026-10-22"}, "expect": card["book"]["fields"]})
    bill_book.forget()

    # ── ядро: участники, режим наших юрлиц, налоговые годы ──────────────
    from app.finance.core import book as core_book

    def customer_of(name: str, number: str) -> str:
        created = owner.post(f"{BASE}/contracts", json=contract_values(name, number))
        assert created.status_code < 400, created.text
        return created.json()["contract"]["values"]["customer"]

    berg = customer_of("ТОО Бергауф Астана", "ОБО/1")
    berg_twin = customer_of("Бергауф Астана", "ОБО/2")
    walk.call(owner, "PUT", "/core/participants/{party_id}/bin", party_id=berg, json={"bin": "260540028486"})
    walk.call(owner, "PATCH", "/core/participants/{party_id}", party_id=berg, json={"full_name": "ТОО «Бергауф Астана»"})
    aliases = walk.call(owner, "POST", "/core/participants/{party_id}/aliases", party_id=berg,
                        json={"name": "BERGAUF ASTANA LLP", "source": "bank"}).json()["aliases"]
    alias_id = next(item["id"] for item in aliases if item["name"] == "BERGAUF ASTANA LLP")
    walk.call(owner, "DELETE", "/core/participants/{party_id}/aliases/{alias_id}", party_id=berg, alias_id=alias_id)
    walk.call(owner, "POST", "/core/participants/{party_id}/bin/reject", party_id=berg_twin, json={"bin": "090540002228"})
    walk.call(owner, "POST", "/core/participants/distinct",
              json={"ids": [customer_of("ТОО Гамма", "ОБО/3"), customer_of("ТОО Дельта", "ОБО/4")]})
    walk.call(owner, "POST", "/core/participants/merge", json={"keep": berg, "drop": berg_twin})
    company = uuid.UUID(owner.get(f"{BASE}/auth/me").json()["company"]["id"])
    core_book.put(company, core_book.build_book(
        [["Наша Фирма", "Заказчик (Название Фирмы)", "По Выписке Банка", "№ Договора"],
         ["BBC", "ТОО Альфа", "ТОО \"АЛЬФА\" БИН 260540028486", "ЮО/141"]],
        title="Сводка", worksheet="Сводка",
    ))
    assert owner.get(f"{BASE}/core/participants-harvest").json()["matched"] == 1
    walk.call(owner, "POST", "/core/participants-harvest")
    core_book.forget()
    walk.call(owner, "POST", "/core/entities/{entity_id}/tax", entity_id=entity,
              json={"valid_from": "2026-01-01", "regime": "general", "vat_payer": True})
    period = owner.get(f"{BASE}/core/participants/{entity}").json()["own"]["tax"][0]["id"]
    walk.call(owner, "DELETE", "/core/entities/{entity_id}/tax/{period_id}", entity_id=entity, period_id=period)
    walk.call(owner, "PUT", "/core/tax-years/{year}", year=2027,
              json={"mrp": "4600", "threshold_mrp": 10000, "vat_rate": "16"})

    # ── ядро: поправка месяца, статус с даты, сверка с книгой ───────────
    def monthly(customer: str, number: str, signed: str) -> str:
        created = owner.post(f"{BASE}/contracts", json=contract_values(customer, number, signed_at=signed, billing="month"))
        assert created.status_code < 400, created.text
        return created.json()["contract"]["id"]

    omega = monthly("ТОО Омегаком", "ОБО/8", "25.05.2026")
    sigma = monthly("ТОО Сигмаком", "ОБО/9", "27.05.2026")
    assert owner.get(f"{BASE}/core/contracts/{omega}/months").status_code == 200
    adjustment = walk.call(owner, "POST", "/core/contracts/{contract_id}/adjustments", contract_id=omega,
                           json={"month": "2026-07", "amount": "5000", "reason": "скидка"}).json()["id"]
    walk.call(owner, "DELETE", "/core/contracts/{contract_id}/adjustments/{document_id}",
              contract_id=omega, document_id=adjustment)
    walk.call(owner, "POST", "/core/contracts/{contract_id}/status-from", contract_id=sigma,
              json={"status_id": paused, "effective_from": "2026-09-01"})
    rows = [
        [str(month), f"{title} 2026", "BBC", customer, "500 000", number, f"01.0{month}.2026", f"{last}.0{month}.2026"]
        for month, title, last in ((6, "ИЮНЬ", 30), (7, "ИЮЛЬ", 31))
        for number, customer in (("ОБО/8", "ТОО Омегаком"), ("ОБО/9", "ТОО Сигмаком"))
    ]
    rows.append(["7", "ИЮЛЬ 2026", "BBC", "ТОО Новичок", "15 000", "БезДог.07-1", "", ""])
    core_book.put(company, core_book.build_book(
        [["Мес", "Техн. 1", "Наша Фирма", "Заказчик (Название Фирмы)", "Сумма Договора", "№ Договора",
          "Период с (нач.)", "Период по (зав.)"], *rows],
        title="Сводка", worksheet="Сводка",
    ))
    view = owner.get(f"{BASE}/core/shadow").json()
    first_day = next(item for item in view["proposals"] if item["kind"] == "period_day" and item["contract_id"] == omega)
    walk.call(owner, "POST", "/core/shadow/apply", json={key: first_day[key] for key in ("contract_id", "kind", "value")})
    walk.call(owner, "POST", "/core/shadow/apply-all", json={"kind": "period_day"})
    walk.call(owner, "POST", "/core/shadow/create", json={"key": view["unknown"][0]["key"]})
    # «Верно у KORT»: у ОБО/9 в июле в таблице другая сумма, у ОБО/8 таблица пропустила 01.07 - 02.07.
    core_book.put(company, core_book.build_book(
        [["Мес", "Техн. 1", "Наша Фирма", "Заказчик (Название Фирмы)", "Сумма Договора", "№ Договора",
          "Период с (нач.)", "Период по (зав.)"],
         ["6", "ИЮНЬ 2026", "BBC", "ТОО Омегаком", "500 000", "ОБО/8", "01.06.2026", "30.06.2026"],
         ["7", "ИЮЛЬ 2026", "BBC", "ТОО Омегаком", "500 000", "ОБО/8", "03.07.2026", "31.07.2026"],
         ["6", "ИЮНЬ 2026", "BBC", "ТОО Сигмаком", "500 000", "ОБО/9", "01.06.2026", "30.06.2026"],
         ["7", "ИЮЛЬ 2026", "BBC", "ТОО Сигмаком", "400 000", "ОБО/9", "01.07.2026", "31.07.2026"]],
        title="Сводка", worksheet="Сводка",
    ))

    def line_of(contract_id: str, month: str) -> dict:
        lines = owner.get(f"{BASE}/core/shadow").json()["lines"]
        return next(line for line in lines if line["contract_id"] == contract_id and line["month"] == month)

    assert line_of(sigma, "2026-07")["codes"] == ["amount"]
    walk.call(owner, "POST", "/core/shadow/verify", json={"contract_id": sigma, "month": "2026-07"})
    walk.call(owner, "DELETE", "/core/shadow/verify/{document_id}", document_id=line_of(sigma, "2026-07")["verified"]["id"])
    assert line_of(omega, "2026-07")["slip"] == "gap"
    slips = walk.call(owner, "POST", "/core/shadow/verify-slips", json={}).json()
    assert slips["done"] == 1, slips
    walk.call(owner, "DELETE", "/core/shadow/verify-batch/{batch}", batch=slips["batch"])
    assert line_of(omega, "2026-07")["verified"] is None
    # «Верно у KORT» вместо предложения «сумма с даты» у ОБО/9.
    amount = next(item for item in owner.get(f"{BASE}/core/shadow").json()["proposals"] if item["contract_id"] == sigma)
    walk.call(owner, "POST", "/core/shadow/verify-proposal", json={"contract_id": sigma, "months": amount["months"]})
    # Оплаты: начало учёта, входящее сальдо из книги и руками, поступление без договора.
    walk.call(owner, "PUT", "/core/payments/settings", json={"since": "2026-06-01"})
    core_book.put(company, core_book.build_book(
        [["Мес", "Техн. 1", "Наша Фирма", "Заказчик (Название Фирмы)", "Сумма Договора", "№ Договора",
          "Период с (нач.)", "Период по (зав.)", "Сальдо Начало (без АВР)"],
         ["6", "ИЮНЬ 2026", "BBC", "ТОО Омегаком", "500 000", "ОБО/8", "01.06.2026", "30.06.2026", "120 000"]],
        title="Сводка", worksheet="Сводка",
    ))
    overview = owner.get(f"{BASE}/core/payments")
    assert overview.status_code == 200, overview.text
    assert overview.json()["opening"]["counts"]["new"] == 1
    walk.call(owner, "POST", "/core/payments/opening")
    walk.call(owner, "PUT", "/core/contracts/{contract_id}/opening", contract_id=sigma, json={"amount": "50000"})
    stray = owner.post(f"{BASE}/operations", json={"kind": "income", "paid_at": "2026-06-15", "amount": "7000",
                                                     "account_to_id": reserve, "comment": "непонятный приход"})
    assert stray.status_code < 400, stray.text
    walk.call(owner, "POST", "/core/payments/{operation_id}/decide", operation_id=stray.json()["id"], json={"action": "none"})
    walk.call(owner, "PUT", "/core/payments/ignore", json={"kind": "payers", "name": "Вознаграждение"})
    # Словарь колонок: «Сумма Договора» компании называется иначе.
    walk.call(owner, "PUT", "/core/book/columns", json={"key": "amount", "name": "Сумма по договору"})
    core_book.forget()
    # Порог НДС: предпросмотр перевода ничего не пишет, перевод - пачка «исполнитель с даты».
    # Договор уже у BBCL - переводим ко второму нашему юрлицу.
    other = owner.post(f"{BASE}/contracts/setup/entities", json={"name": "BBC Astana", "code": "BBCAS"})
    assert other.status_code < 400, other.text
    move = {"year": 2026, "moves": [{"contract_id": omega, "target": other.json()["id"], "since": "2026-12-01"}]}
    assert owner.get(f"{BASE}/core/vat", params={"year": 2026}).status_code == 200
    assert owner.post(f"{BASE}/core/vat/preview", json=move).status_code == 200
    walk.call(owner, "POST", "/core/vat/transfer", json=move)
    # Переход: счета BBCL - из KORT, и снова из таблицы.
    walk.call(owner, "PUT", "/core/switch/{link}/{entity_id}", link="bills", entity_id=entity, json={"source": "kort"})
    assert owner.get(f"{BASE}/bills").json()["source"]["kort"] == ["BBCL"]
    back = owner.put(f"{BASE}/core/switch/bills/{entity}", json={"source": "book"})
    assert back.status_code == 200 and back.json()["sources"]["bills"] == {}, back.text
    # Слова компании: слово колонки «Счет» со своим смыслом.
    walk.call(owner, "PUT", "/core/table-words", json={"part": "bill", "word": "Сегодня день выставления", "meaning": "soon"})
    # Первые шаги: юрлицо по новому коду, договор по строкам таблицы, сальдо на начало учёта.
    core_book.put(company, core_book.build_book(
        [["Мес", "Техн. 1", "Наша Фирма", "Заказчик (Название Фирмы)", "Сумма Договора", "№ Договора",
          "Период с (нач.)", "Период по (зав.)", "Сальдо Начало (без АВР)"],
         ["6", "ИЮНЬ 2026", "СТАРТ", "ТОО Стартап", "40 000", "СТ/1", "01.06.2026", "30.06.2026", "15 000"]],
        title="Сводка", worksheet="Сводка",
    ))
    assert owner.get(f"{BASE}/core/start").json()["contracts"]["missing"] == 1
    walk.call(owner, "POST", "/core/start/{step}", step="entities")
    walk.call(owner, "POST", "/core/start/{step}", step="contracts")
    walk.call(owner, "POST", "/core/start/{step}", step="since")
    core_book.forget()

    walk.call(owner, "DELETE", "/contracts/setup/summary")

    # ── загрузка реестра ────────────────────────────────────────────────
    upload = walk.call(owner, "POST", "/contracts/imports",
                       files={"file": ("реестр.xlsx", _registry_file(), XLSX)}).json()
    columns = {
        column["id"]: {"action": "note"}
        for block in _section(upload, "columns")["items"]
        for column in block["columns"]
        if column["decision"].get("action") == "ask"
    }
    own = [item["key"] for item in _section(upload, "entities")["items"] if item["own"]]
    decided = owner.post(
        f"{BASE}/contracts/imports/{upload['id']}/decide",
        json={"decisions": {"columns": columns, "entities": {"confirmed": True, "own": own}}},
    )
    assert decided.status_code == 200, decided.text
    walk.call(owner, "POST", "/contracts/imports/{batch_id}/apply", batch_id=upload["id"])
    again = owner.post(f"{BASE}/contracts/imports", files={"file": ("реестр.xlsx", _registry_file(), XLSX)}).json()
    walk.call(owner, "POST", "/contracts/imports/{batch_id}/cancel", batch_id=again["id"])

    # ── передача владения - последней: дальше владелец уже администратор ──
    heir = owner.post(
        f"{BASE}/auth/members",
        json={"email": "heir@bbc.kz", "password": "heir-pass-1", "role": "admin", "full_name": "Наследник Дел"},
    )
    assert heir.status_code < 400, heir.text
    walk.call(owner, "POST", "/auth/owner", json={"user_id": heir.json()["id"], "password": "new-pass-123"})

    # ── все изменяющие маршруты пройдены или названы с причиной ─────────
    mutating = {
        (method, full[len(BASE):])
        for full, route in api_routes()
        if full.startswith(BASE)
        for method in route.methods
        if method in ("POST", "PATCH", "PUT", "DELETE")
    }
    # Пустой обход - не «всё покрыто», а слепой тест (FastAPI 0.141).
    assert mutating, "обход не нашёл ни одного изменяющего маршрута"
    missing = sorted(mutating - walk.covered - set(ALLOWED))
    assert not missing, "изменяющие маршруты без проверенного события: " + ", ".join(f"{m} {p}" for m, p in missing)
    stale = sorted(set(ALLOWED) - mutating)
    assert not stale, f"в ALLOWED маршруты, которых больше нет: {stale}"
    assert not (walk.covered & set(ALLOWED)), "маршрут и пишет событие, и стоит в ALLOWED - уберите из списка"


def test_zhurnal_nastroyki_govorit_chto_bylo_i_chto_stalo(app: FastAPI) -> None:
    """28.09.2026 владелец случайно спрятал «Планируемый срок завершения», а
    журнал записал одно «поле изменено» и ключ - ни что сделано, ни как было.
    Теперь запись говорит словами и хранит «было → стало» того, что поменялось."""
    owner = client(app)
    owner.post(
        f"{BASE}/auth/register",
        json={"email": "setup@bbc.kz", "password": PASSWORD, "company": "BBC", "full_name": "Ермеков Нурболат"},
    )
    owner.get(f"{BASE}/contracts/schema")

    def last(kind: str) -> ActionLog:
        with finance_session() as session:
            row = session.scalars(sa.select(ActionLog).where(ActionLog.kind == kind).order_by(ActionLog.at.desc())).first()
            session.expunge(row)
            return row

    assert owner.patch(f"{BASE}/contracts/setup/fields/planned_end_at", json={"hidden": True}).status_code == 200
    hidden = last("contracts.setup.field_update")
    assert "спрятано" in hidden.title and "Планируемый срок завершения" in hidden.title, hidden.title
    assert hidden.before == {"key": "planned_end_at", "hidden": False}
    assert hidden.after == {"key": "planned_end_at", "hidden": True}

    # Вернули - запись тоже говорит, что поле снова видно.
    assert owner.patch(f"{BASE}/contracts/setup/fields/planned_end_at", json={"hidden": False}).status_code == 200
    back = last("contracts.setup.field_update")
    assert "снова в листе и карточке" in back.title, back.title

    status = owner.post(f"{BASE}/contracts/setup/lists/status", json={"value": "На паузе"}).json()["id"]
    assert owner.patch(f"{BASE}/contracts/setup/values/{status}", json={"value": "Ждёт подписи клиента"}).status_code == 200
    renamed = last("contracts.setup.value_update")
    assert renamed.before["value"] == "На паузе" and renamed.after["value"] == "Ждёт подписи клиента", renamed.before
    assert "переименовано" in renamed.title
