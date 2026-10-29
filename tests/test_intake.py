"""Приёмник: что значит строка листа, шапка в две строки, служебные и скрытые строки.

Файлы - полигона «Тумар» (`tests/polygon`): компания с 1С, Excel и выписками,
которую KORT обязан понять без ручной разметки.
"""
from __future__ import annotations

import io

import openpyxl
import pytest
from polygon import sources
from polygon.world import build

from app.finance.intake import forms, header, recognize, rows, tables


@pytest.fixture(scope="module")
def files() -> dict:
    return sources.render(build())


def _sheets(name: str, data: bytes) -> dict[str, recognize.Sheet]:
    return {sheet.name: sheet for sheet in recognize.recognize_file(tables.read(data, name), name)}


def test_kazhdyi_fail_poligona_uznaet_svoyu_formu(files: dict) -> None:
    expected = {
        "1С Тумар Сервис/ОСВ 1210 июнь 2026.xlsx": forms.BALANCE,
        "1С Тумар Офис/Контрагенты.xlsx": forms.PARTICIPANT,
        "1С Тумар Офис/Реализация товаров и услуг 2026.xlsx": forms.DOCUMENT,
        "1С Тумар Офис/Счета на оплату 2026.xlsx": forms.DOCUMENT,
        "Excel/Реестр договоров 2026.xlsx": forms.CONTRACT,
        "Excel/Аренда - начисления 2026.xlsx": forms.CONTRACT_MONTH,
        "Банк/Выписка Kaspi Тумар Сервис 2026-06.xlsx": forms.PAYMENT,
    }
    for name, form in expected.items():
        sheets = _sheets(name, files[name])
        assert sheets and all(sheet.form == form for sheet in sheets.values()), (name, [(s.name, s.form) for s in sheets.values()])


def test_shapka_v_dve_stroki_skladyvaetsya(files: dict) -> None:
    """Ф1: «Договор» объединён над «№ | дата» - это номер и дата договора; ОСВ - «Сальдо» над «Дебет | Кредит»."""
    registry = _sheets("Реестр.xlsx", files["Excel/Реестр договоров 2026.xlsx"])["Клининг"]
    keys = {column.key for column in registry.reading.columns}
    assert {"number", "signed_at", "customer", "party_bin", "status"} <= keys
    assert registry.title == "РЕЕСТР ДОГОВОРОВ: КЛИНИНГ"
    osv = _sheets("ОСВ.xlsx", files["1С Тумар Офис/ОСВ 1210 июнь 2026.xlsx"])["TDSheet"]
    assert {"opening_debit", "opening_credit", "closing_debit", "closing_credit"} <= {c.key for c in osv.reading.columns}


def test_skrytaya_stroka_voprosom_a_itog_propuskaetsya(files: dict) -> None:
    """Ф9 и Ф6: скрытый черновик - вопрос «пропустить или завести»; строка «Итого» - не данные."""
    sheet = _sheets("Реестр.xlsx", files["Excel/Реестр договоров 2026.xlsx"])["Клининг"]
    assert len(sheet.hidden) == 1
    assert [q["kind"] for q in sheet.questions] == ["hidden"]
    assert sheet.service.get("total") == 1


def test_dve_kolonki_summa_otkaz(files: dict) -> None:
    """Ф5: в листе «Октябрь» две колонки «Сумма» - лист не читается, а не берётся первая."""
    sheets = _sheets("Аренда.xlsx", files["Excel/Аренда - начисления 2026.xlsx"])
    assert [r["kind"] for r in sheets["Октябрь"].refusals] == ["duplicate"]
    assert not sheets["Июль"].refusals  # Ф4: колонки переставлены - читается по названиям
    assert sheets["Июль"].reading.at("accrued") is not None


def test_neponyatnyi_list_vopros_a_ne_samaya_pohozhaya_forma() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Фрукт", "Цвет", "Вкус"])
    sheet.append(["Яблоко", "Красный", "Сладкий"])
    buffer = io.BytesIO()
    book.save(buffer)
    found = recognize.recognize_file(tables.read(buffer.getvalue(), "x.xlsx"), "x.xlsx")
    assert found[0].form is None and found[0].questions[0]["kind"] == "form"


def test_shapka_bez_edinicy_i_skobok() -> None:
    assert header.core("Сумма в месяц, тг") == "сумма в месяц"
    assert header.core("Стоимость (тенге)") == "стоимость"
    assert header.core("Отдел (ОБО, НО)") == "отдел"


def test_itog_restra_strozhe_vypiski() -> None:
    assert rows.is_total(["", "Итого:", None, 150000])
    assert rows.is_numbering([1, 2, 3, 4])
    assert not rows.is_numbering([1, 3, 4])


def test_slovo_itoga_v_primechanii_ne_delaet_dogovor_itogom() -> None:
    """Сводная BBC, НО/113: примечание «в итоге закрыли как разовую» - строка договора, а не итог.

    Слово итога где угодно в строке годится выписке («Остаток на конец» без
    даты), но у записей итог - только подпись единственным текстом строки:
    договор с примечанием пропадал из «Сравнить», реестра и сводной файлом молча.
    """
    from app.finance.intake import layout

    assert not rows.is_label_total(["ТОО Альфа", "№1", "в итоге закрыли как разовую", 10000])
    assert not rows.is_label_total(["ТОО Баланс", "Д-1", "всего три платежа", 5000])
    assert not rows.is_label_total(["Итоговый Консалтинг", "123456789012"])
    assert rows.is_label_total(["", "Итого:", None, 150000, "#N/A", "-"])
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Месяц", "Наше ТОО", "Клиент", "№ договора", "Сумма", "Примечание"])
    sheet.append(["ИЮЛЬ 2026", "СВ", "ТОО Альфа", "А-1", 10000, "в итоге закрыли как разовую"])
    sheet.append(["ИЮЛЬ 2026", "СВ", "ТОО Баланс", "Б-2", 5000, "всего три платежа"])
    sheet.append(["Итого", None, None, None, 15000, None])
    buffer = io.BytesIO()
    book.save(buffer)
    [table] = tables.read(buffer.getvalue(), "x.xlsx")
    [found] = recognize.recognize_file([table], "x.xlsx")
    assert found.form == forms.CONTRACT_MONTH and found.rows == 2 and found.service == {"total": 1}
    read = layout.read(table, found.reading.form)
    assert [row.values["customer"] for row in read.rows] == ["ТОО Альфа", "ТОО Баланс"]


def test_skrytye_stroki_i_obedineniya_chitayutsya() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Договор", None])
    sheet.append(["№", "дата"])
    sheet.append(["Д-1", "01.01.2026"])
    sheet.merge_cells("A1:B1")
    sheet.row_dimensions[3].hidden = True
    buffer = io.BytesIO()
    book.save(buffer)
    [table] = tables.read(buffer.getvalue(), "x.xlsx")
    assert table.hidden == frozenset({2})
    assert table.spans(0) == [(0, 1)]
    fast = tables.read(buffer.getvalue(), "x.xlsx", detail=False)[0]
    assert fast.rows[:3] == table.rows[:3] and not fast.hidden


def test_opisanie_faila_dlya_marshruta(files: dict) -> None:
    """`POST /core/intake` отвечает формой каждого листа; PDF - выпиской, не разбирая."""
    out = recognize.describe(files["Excel/Аренда - начисления 2026.xlsx"], "Аренда.xlsx")
    assert {sheet["form"] for sheet in out["sheets"]} == {forms.CONTRACT_MONTH}
    october = next(sheet for sheet in out["sheets"] if sheet["sheet"] == "Октябрь")
    assert october["refusals"] and october["header_line"] == 3
    assert recognize.describe(b"%PDF-1.4 x", "выписка.pdf")["sheets"][0]["form"] == forms.PAYMENT
    assert recognize.describe(b"\x00\x01garbage", "x.bin")["error"]
